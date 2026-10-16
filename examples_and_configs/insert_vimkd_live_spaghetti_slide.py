"""
insert_vimkd_live_spaghetti_slide.py

PI deck slide 4 (FF 2026-10-04): live VimKD centrosome-synapse distance per cell over time,
siCtrl | siVim side by side on a shared y-axis (Supplementary Figure 6.2 D/E restyled by
fixed_cells_general_20240801_vim_paper/vimkd_centrosome_by_day_20261004/gen_vimkd_live_spaghetti_shared_y.m).
Same style as insert_vimkd_centrosome_slide.py (16:9, 48 pt bold title, no footer).

Usage:
    conda run -n PPT_editing python examples_and_configs/insert_vimkd_live_spaghetti_slide.py
"""

import os
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import insert_channel_misalignment_slides as CM

SRC = Path("K:/FF/Papers/vimentin_paper/Figures/postprocessing_output/Figure5/VimKD_live_cell/"
           "spaghetti_by_condition_shared_y_20261004")
PANELS = [SRC / f"{c}_distance_to_synapse_rainbow_y0_16.png" for c in ("siCtrl", "siVim")]
OUT_DIR = Path("K:/FF/PPT/PPT_autogeneration/Fixed Jurkats, Miscellaneous/VimentinKD")
OUT = OUT_DIR / "VimKD_live_centrosome_synapse_distance_20261004.pptx"
TITLE = "Live cells: centrosome-synapse distance over time"   # draft; FF to edit
TITLE_PT, TITLE_H, GAP = 48, 1.65, 0.4
LABELS = ("siCtrl", "siVim")
LABEL_PT, LABEL_H, LABEL_W, LABEL_INSET = 24, 0.5, 1.6, 0.12
LABEL_RIGHT = 0.25   # clear of the right axis line and its ticks


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(CM.W), Inches(CM.H)
    s = CM.M._new_slide(prs)
    CM.M.add_textbox(s, TITLE, CM.MARGIN, CM.TITLE_TOP, CM.IMG_W, TITLE_H,
                     font_pt=TITLE_PT, color=CM.BLACK, bold=True)
    s.shapes[-1].text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    src = OUT_DIR / "source_images"
    src.mkdir(parents=True, exist_ok=True)
    imgs, aspects = [], []
    for p in PANELS:
        out = src / p.name
        CM.trim(Image.open(p)).save(out)
        with Image.open(out) as im:
            aspects.append(im.width / im.height)
        imgs.append(out)
    # condition names as slide text INSIDE each panel's top-right corner, where the figure used to
    # draw them (FF 2026-10-04: in the PPT, not in the images). Placed after the picture so it sits on top.
    img_top = CM.TITLE_TOP + TITLE_H + 0.08
    h = min(CM.H - img_top - CM.MARGIN, (CM.IMG_W - GAP) / sum(aspects))
    widths = [h * a for a in aspects]
    left = CM.MARGIN + (CM.IMG_W - sum(widths) - GAP) / 2
    for img, w, lab in zip(imgs, widths, LABELS):
        s.shapes.add_picture(str(img), Inches(left), Inches(img_top), Inches(w), Inches(h))
        CM.M.add_textbox(s, lab, left + w - LABEL_W - LABEL_RIGHT, img_top + LABEL_INSET, LABEL_W, LABEL_H,
                         font_pt=LABEL_PT, color=CM.BLACK, bold=True, align=PP_ALIGN.RIGHT)
        left += w + GAP
    if OUT.exists():
        CM.M.backup_presentation(str(OUT), backup_base=str(OUT_DIR / "backups"))
    prs.save(str(OUT))
    print(f"Done -> {OUT}")


if __name__ == "__main__":
    main()
