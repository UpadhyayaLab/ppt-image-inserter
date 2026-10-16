"""
make_naiveCTL_GzmB_4up_template.py

Build NaiveCTL_GzmB_histone_summary_template.pptx: a copy of
NaiveCTL_GzmB_bTub_summary_template.pptx (slides 0-3 unchanged) with one extra
layout slide, slide 4 = 4-up square panels in a single row. Used by
gen_naiveCTL_GzmB_bTub_summary_config.py (--dataset 10152025) for the
glass vs glass 1-16 comparison at 6 / 24 / 48 / 72 h.

Placeholder pictures reuse slide 3's picture blob; batch_insert_images.py
(auto_position) only reads their positions and sizes.

Usage:
    conda run -n PPT_editing python examples_and_configs/make_naiveCTL_GzmB_4up_template.py
"""

import io
import os
import shutil

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

SRC = "K:/FF/PPT/PPT_autogeneration/Naive_CTL/NaiveCTL_GzmB_bTub_summary_template.pptx"
DST = "K:/FF/PPT/PPT_autogeneration/Naive_CTL/NaiveCTL_GzmB_histone_summary_template.pptx"

N_PANELS = 4
PANEL_IN = 3.2      # square panel side
GAP_IN = 0.1
# Vertically centered on slide 3's 3-up row (top 2.3 in, 4.1 in panels).
ROW_CENTER_IN = 2.3 + 4.1 / 2


def main():
    if os.path.exists(DST):
        raise SystemExit(f"{DST} exists; delete it first to rebuild.")
    shutil.copyfile(SRC, DST)
    prs = Presentation(DST)
    if len(prs.slides) != 4:
        raise SystemExit(f"Expected 4 template slides in {SRC}, found {len(prs.slides)}.")

    ref = prs.slides[3]
    blob = next(sh for sh in ref.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE).image.blob
    slide = prs.slides.add_slide(ref.slide_layout)
    # Drop the layout's empty placeholders so only pictures remain (as on slide 3).
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)

    slide_w_in = prs.slide_width / 914400.0
    row_w = N_PANELS * PANEL_IN + (N_PANELS - 1) * GAP_IN
    left0 = (slide_w_in - row_w) / 2
    top = ROW_CENTER_IN - PANEL_IN / 2
    for k in range(N_PANELS):
        slide.shapes.add_picture(io.BytesIO(blob), Inches(left0 + k * (PANEL_IN + GAP_IN)),
                                 Inches(top), Inches(PANEL_IN), Inches(PANEL_IN))
    prs.save(DST)
    print(f"Wrote {DST} ({len(prs.slides)} slides; slide 4 = {N_PANELS}-up, "
          f"{PANEL_IN} in panels at top {top:.2f} in)")


if __name__ == "__main__":
    main()
