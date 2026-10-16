"""
make_ctl_em_deck_template.py

Build CTL_EM_centrosome_invagination_template.pptx (16:9) for the OpenOrganelle CTL FIB-SEM deck
(FF 2026-10-02 advisor progress report: Ritter et al. 2022 Fig. 2A cell jrc_ctl-id8-2 vs jrc_ctl-id8-4):
  slide 0 = title slide
  slide 1 = one large picture placeholder under the title band (batch_insert_images.py auto_position
            fits each image inside it, aspect preserved, centered)

Placeholder picture = a plain white PNG made here; only its position and size are read.
Filled by configs/fixed/CTLs/config_CTL_EM_centrosome_invagination.yaml.

Usage:
    conda run -n PPT_editing python examples_and_configs/make_ctl_em_deck_template.py
"""

import io
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

DST = "K:/FF/PPT/PPT_autogeneration/CTL_EM/CTL_EM_centrosome_invagination_template.pptx"
BLANK = "K:/FF/PPT/PPT_autogeneration/CTL_EM/blank_for_screenshot.png"

SLIDE_W, SLIDE_H = 13.333, 7.5
# picture box: below a one-line title (32 pt), thin side margins; config titles are kept to one line
PIC = dict(left=0.15, top=0.8, width=SLIDE_W - 0.3, height=SLIDE_H - 0.8 - 0.45)  # under a ONE-line title, above the footer
# citation footer on the template slide, so every inserted slide carries it (FF 2026-10-02 "add citations");
# references as printed in the Ritter et al. supplement (science.abl3855_sm.pdf, header and ref. 13)
CITE = "Data: Ritter et al., Science 376, 377 (2022) · OpenOrganelle: Xu et al., Nature 599, 147 (2021)"

TITLE = "Centrosome and nuclear invagination in killer T cells"
SUBTITLE = "FIB-SEM of CTLs attacking tumor cells (Ritter et al. 2022)"
FONT = "Calibri"     # FF 2026-10-02 "use established Calibri"; set explicitly, not left to the theme


def main():
    os.makedirs(os.path.dirname(DST), exist_ok=True)
    if os.path.exists(DST):
        raise SystemExit(f"{DST} exists; delete it first to rebuild.")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)

    s0 = prs.slides.add_slide(prs.slide_layouts[0])          # title layout
    # the default layouts are placed for a 10 in slide: re-center both placeholders on the 13.333 in slide
    for ph, top, hgt in ((s0.shapes.title, 2.0, 1.6), (s0.placeholders[1], 3.8, 1.4)):
        ph.left, ph.top, ph.width, ph.height = Inches(0.8), Inches(top), Inches(SLIDE_W - 1.6), Inches(hgt)
    s0.shapes.title.text = TITLE
    s0.shapes.title.text_frame.paragraphs[0].font.size = Pt(44)
    s0.shapes.title.text_frame.paragraphs[0].font.name = FONT
    s0.placeholders[1].text = SUBTITLE
    for p in s0.placeholders[1].text_frame.paragraphs:
        p.font.size = Pt(28)
        p.font.name = FONT

    s1 = prs.slides.add_slide(prs.slide_layouts[6])          # blank layout
    buf = io.BytesIO()
    Image.new("RGB", (16, 9), "white").save(buf, format="PNG")
    buf.seek(0)
    s1.shapes.add_picture(buf, Inches(PIC["left"]), Inches(PIC["top"]), Inches(PIC["width"]), Inches(PIC["height"]))
    tb = s1.shapes.add_textbox(Inches(0.3), Inches(SLIDE_H - 0.42), Inches(SLIDE_W - 0.6), Inches(0.35))
    tb.text_frame.text = CITE
    para = tb.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    para.font.size, para.font.name = Pt(16), FONT
    para.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    prs.save(DST)
    # blank white placeholder with the picture box's aspect: slides for FF's own screenshots of
    # Ritter et al. 2022 Fig. 2A and Fig. S2B, pasted by FF (FF 2026-10-02 "I'll add screenshots. leave room")
    w = 1600
    Image.new("RGB", (w, round(w * PIC["height"] / PIC["width"])), "white").save(BLANK)
    print(f"Wrote {DST}: slide 0 title, slide 1 picture box {PIC}")


if __name__ == "__main__":
    main()
