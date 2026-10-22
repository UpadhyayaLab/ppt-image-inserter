#!/usr/bin/env python3
r"""build_pmlc_results_deck.py  Comprehensive pMLC results deck: every key pMLC
metric across ALL 3 replicates (Mar 12 / May 10 / Jun 07, 2026) and ALL
timepoints, with CAT-vs-FMC-on-one-graph views. Pulls the EXISTING compiled
grid panels (one tile per replicate date) -- no recompile, no new plotting.

Sources (read-only):
  compiled_pmlc_v2_20260831\grid_panels\<group>\<metric>_grid.png   (per-metric grids)
  combined_overlays_20260913\combined_{profile,gradient}_<cond>_<date>_.png

One slide per metric = up to four comparison-group grids (each already CAT+FMC or
timecourse) in a 2x2, labelled by group. Then a combined actin+pMLC overlay
appendix (one slide per timepoint, CAT row over FMC row, 3 replicate columns).

Layout helpers copied from the vimentin Fig4 builder
(fixed_cells_general_20240801_vim_paper\fig4_lata_consistency_20260902\
 build_fig4_metrics_deck.py): fit_in_box aspect-fits, never stretches.

Output: K:\FF\PPT\PPT_autogeneration\CART\pMLC\CART_pMLC_results_all_replicates_<yyyymmdd>.pptx
Run:    python build_pmlc_results_deck.py
"""
from __future__ import annotations

import datetime
import os
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

NAS = Path(r'\\10.225.56.176\ULab_data\User_data\Kiet\results_compiled\pmlc_joint')
GRIDS = NAS / 'compiled_pmlc_v2_20260831' / 'grid_panels'
OVERLAYS = NAS / 'combined_overlays_20260913'
OUT_DIR = Path(r'K:\FF\PPT\PPT_autogeneration\CART\pMLC')
OUT = OUT_DIR / f'CART_pMLC_results_all_replicates_{datetime.date.today():%Y%m%d}.pptx'

SLIDE_W, SLIDE_H = 13.333, 7.5           # 16:9 review deck
BLACK = RGBColor(0, 0, 0); GREY = RGBColor(0x55, 0x55, 0x55)
MU = '\u00b5'

# Metric -> display title. Region-norm metrics carry the batch-safe claim; the
# profile/gradient show the peripheral pMLC ring; MFI + inner-outer are the
# draft's featured pMLC panels.
METRICS = [
    ('pMLC_synapse_peripheral_MFI_norm',      'pMLC Peripheral MFI / Synapse Mean'),
    ('pMLC_synapse_central_MFI_norm',         'pMLC Central MFI / Synapse Mean'),
    ('pMLC_synapse_distal_MFI_norm',          'pMLC Distal MFI / Synapse Mean'),
    ('pMLC_synapse_ray_profile_norm',         'pMLC Radial Profile (/ Synapse Mean)'),
    ('pMLC_synapse_ray_profile_norm_grad',    'pMLC Radial Gradient'),
    ('pMLC_synapse_inner_outer_ratio_70pct',  'pMLC Inner-Outer Ratio'),
    ('pMLC_synapse_MFI',                      'pMLC Synapse MFI'),
]
# Comparison groups that put CAT & FMC together / show the timecourse. Order =
# panel order on the slide (2x2). Profile fields only exist in the first three.
GROUPS = [
    ('CAT_vs_FMC_by_timepoint', 'CAT vs FMC at Each Timepoint'),
    ('time_course_by_cell_type', 'Time Course within Each Cell Type'),
    ('timecourse_line_graphs',   'Timecourse (Mean \u00b1 CI, CAT vs FMC)'),
    ('timecourse_scatter_plots', 'Timecourse (Scatter, CAT vs FMC)'),
]
DATES = ['Mar_12', 'May_10', 'Jun_07']
TIMEPOINTS = ['5min', '10min', '15min']


def add_text(slide, left, top, width, height, text, size, *, bold=False,
             color=BLACK, align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = wrap
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        p.space_after = Pt(3)
    return tb


def fit_in_box(slide, png, left, top, box_w, box_h):
    """Aspect-fit a PNG inside a box, centered. Returns True if placed."""
    png = Path(png)
    if not png.is_file():
        return False
    with Image.open(png) as im:
        iw, ih = im.size
    ar = iw / ih
    if ar >= box_w / box_h:
        w = box_w; h = w / ar
    else:
        h = box_h; w = h * ar
    slide.shapes.add_picture(str(png), Inches(left + (box_w - w) / 2),
                             Inches(top + (box_h - h) / 2), width=Inches(w))
    return True


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def text_to_front(slide):
    tree = slide.shapes._spTree
    for sh in list(slide.shapes):
        if sh.has_text_frame:
            el = sh._element; tree.remove(el); tree.append(el)


def metric_slide(prs, field, title):
    """One slide: up to four comparison-group grids for this metric in a 2x2."""
    grids = [(g, lbl, GRIDS / g / f'{field}_grid.png') for g, lbl in GROUPS]
    grids = [(g, lbl, p) for g, lbl, p in grids if p.is_file()]
    if not grids:
        print(f'  [skip] {field}: no grids on disk')
        return
    s = blank(prs)
    add_text(s, 0.4, 0.12, SLIDE_W - 0.8, 0.6, title, 26, bold=True, align=PP_ALIGN.CENTER)
    # 2x2 cells.
    cell_w = (SLIDE_W - 0.6) / 2
    cell_h = (SLIDE_H - 0.9) / 2
    positions = [(0.3, 0.8), (0.3 + cell_w, 0.8),
                 (0.3, 0.8 + cell_h), (0.3 + cell_w, 0.8 + cell_h)]
    for (g, lbl, p), (lx, ly) in zip(grids, positions):
        add_text(s, lx + 0.1, ly, cell_w - 0.2, 0.3, lbl, 12, bold=True, color=GREY)
        fit_in_box(s, p, lx + 0.05, ly + 0.32, cell_w - 0.1, cell_h - 0.4)
    print(f'  [ok] {field}: {len(grids)} grid(s)')


def overlay_appendix(prs):
    """One slide per timepoint: combined actin+pMLC radial profile, CAT row over
    FMC row, three replicate columns."""
    for tp in TIMEPOINTS:
        # need at least one panel to bother
        any_here = any((OVERLAYS / f'combined_profile_{cc}_{tp}_{d}_.png').is_file()
                       for cc in ('CAT', 'FMC') for d in DATES)
        if not any_here:
            continue
        s = blank(prs)
        add_text(s, 0.4, 0.12, SLIDE_W - 0.8, 0.6,
                 f'Combined Actin (green) + pMLC (magenta) Radial Profile \u2014 {tp}',
                 22, bold=True, align=PP_ALIGN.CENTER)
        cell_w = (SLIDE_W - 0.6) / 3
        cell_h = (SLIDE_H - 1.1) / 2
        for row, cc in enumerate(('CAT', 'FMC')):
            add_text(s, 0.15, 0.85 + row * cell_h, 0.9, 0.3, cc, 16, bold=True, color=GREY)
            for col, d in enumerate(DATES):
                p = OVERLAYS / f'combined_profile_{cc}_{tp}_{d}_.png'
                fit_in_box(s, p, 0.3 + col * cell_w, 0.85 + row * cell_h,
                           cell_w - 0.1, cell_h - 0.15)
        print(f'  [ok] overlay appendix {tp}')


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W); prs.slide_height = Inches(SLIDE_H)

    # Title slide.
    s = blank(prs)
    add_text(s, 0.6, 2.6, SLIDE_W - 1.2, 1.2,
             'pMLC at the CAR-T Synapse: All Replicates, All Timepoints',
             30, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, 0.6, 3.9, SLIDE_W - 1.2, 1.6, [
        'CAT (low-affinity, blue) vs FMC63 (high-affinity, red). Three biological replicates '
        '(Mar 12, May 10, Jun 07 2026), 6 conditions each (CAT/FMC x 5/10/15 min).',
        'Panels are the compiled per-replicate grids (one tile per date) from '
        'compiled_pmlc_v2_20260831; region metrics normalized to each cell\u2019s synapse mean. '
        'Replicates shown separately (not pooled).',
        f'Built {datetime.date.today():%Y-%m-%d} from '
        r'\\10.225.56.176\ULab_data\User_data\Kiet\results_compiled\pmlc_joint.'],
        13, color=GREY, align=PP_ALIGN.CENTER)

    for field, title in METRICS:
        metric_slide(prs, field, title)
    overlay_appendix(prs)

    # Pointer slide.
    s = blank(prs)
    add_text(s, 0.5, 0.3, SLIDE_W - 1.0, 0.5, 'Where everything lives', 22, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, 0.5, 1.0, SLIDE_W - 1.0, 5.5, [
        f'Per-metric grids (this deck): {GRIDS}',
        f'Combined actin+pMLC overlays: {OVERLAYS}',
        'Metrics: Rey-Suarez et al. 2021 region analysis (central r<0.4 / peripheral 0.4-0.8 / '
        'distal r>0.8), all normalized to each cell\u2019s synapse-mean intensity.',
        'Config: config\\compilation\\d20260312_20260510_20260607_pmlc_compilation_config.m; '
        'compile output compiled_pmlc_v2_20260831.'],
        12, color=GREY)

    for sl in prs.slides:
        text_to_front(sl)
    prs.save(str(OUT))
    print(f'\nsaved {OUT} ({len(prs.slides)} slides)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
