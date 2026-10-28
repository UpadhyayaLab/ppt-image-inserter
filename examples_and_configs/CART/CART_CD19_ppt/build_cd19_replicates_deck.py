#!/usr/bin/env python3
r"""build_cd19_replicates_deck.py  CD19 coculture replicate deck: the paper's Figure 6
metrics, experiment by experiment, and whether the experiments agree.

Reads (never plots, never tests) the folder written by volscope's
scripts/compare_experiments.py:
  <replicates>\spec.json                          experiments + metric list (the paper's)
  <replicates>\consistency.csv                    per metric x timepoint, per experiment:
                                                  n, medians, Cliff's delta, BH p, verdict
  <replicates>\<experiment>\<metric>_by_timepoint.png   one panel per experiment
and each experiment's comparison\manifest.json for the metric labels.

Slides: title; one slide per metric with the experiments side by side (shared y, block labels
only); last, the summary table (verdict per metric x timepoint, each experiment's stars).
No notes slide (FF 2026-10-25: "not useful"). Review-deck rules: paper-figures skill section 8.

Output: K:\FF\PPT\PPT_autogeneration\CART\CD19\CART_CD19_analysis_summary_<yyyymmdd>.pptx
(FF 2026-10-27: "call it analysis_summary not replicates"; was CART_CD19_replicates_*)
Run (PowerShell, base miniforge -- has python-pptx + pywin32):
  & 'C:\Users\Upadhyaya_Lab\miniforge3\python.exe' build_cd19_replicates_deck.py `
      --replicates '\\10.225.56.176\ULab_data\User_data\Kiet\volscope_replicates_03102026_20261001' --export
"""
from __future__ import annotations

import argparse
import csv
import datetime
import json
import math
import sys
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'CART_pMLC_ppt'))
import build_pmlc_profiles_deck as bpd  # noqa: E402

DECK_DIR = Path(r'K:\FF\PPT\PPT_autogeneration\CART\CD19')
REPLICATES = Path(r'\\10.225.56.176\ULab_data\User_data\Kiet\volscope_replicates_03102026_20261001')

VERDICT_COLOR = {                           # summary-table fill per verdict
    'replicated': RGBColor(0xC6, 0xEF, 0xCE),
    'same direction': RGBColor(0xFF, 0xEB, 0x9C),
    # 'sig. in <exp> only': significant in one, the n.s. one pointing the other way. The
    # MATLAB (Bug hunt, 2026-10-26) treats a n.s. direction as a null effect, so this is
    # significant-in-one like 'same direction', not a conflict; it fell through to the white
    # 'missing' fill before this entry.
    'sig. in': RGBColor(0xFF, 0xEB, 0x9C),
    'n.s. in all': RGBColor(0xEE, 0xEE, 0xEE),
    'conflict': RGBColor(0xFF, 0xC7, 0xCE),
    'missing': RGBColor(0xFF, 0xFF, 0xFF),
}
# What each color means, for the key under the summary table (FF 2026-10-25). Mirrors
# replicate_verdict in matlab/volscope_compare_experiments.m: since 2026-10-26 a conflict
# needs BOTH experiments significant.
VERDICT_KEY = [
    ('replicated', 'replicated: significant in both, same direction'),
    ('same direction', 'significant in one experiment only'),
    ('conflict', 'conflict: significant in both, opposite directions'),
    ('n.s. in all', 'n.s. in both experiments'),
]


def experiment_label(name):
    """'03102026' (MMDDYYYY) or '20261001' (YYYYMMDD) -> 'Mar 10, 2026'; else as is."""
    for fmt in ('%Y%m%d', '%m%d%Y'):
        try:
            d = datetime.datetime.strptime(name, fmt)
        except ValueError:
            continue
        if 2020 <= d.year <= 2035:
            return f'{d:%b} {d.day}, {d.year}'
    return name


def stars(p):
    """volscope's tiers, which top out at *** (sigstar_within_ylim.m / stars_for_p in
    volscope's matlab/compare): the panels on the slides are volscope figures, so the
    table must print the glyph they show. bpd.stars is CART_fixed's, which adds ****
    (p <= 1e-4) and printed **** in the table over a *** panel."""
    p = float(p)
    if math.isnan(p) or p > 0.05:
        return 'n.s.'
    return '***' if p <= 1e-3 else '**' if p <= 1e-2 else '*'


def slide_order(spec):
    """The spec's metric order (paper metrics, then extras); older specs: violins, fractions."""
    return spec.get('order') or spec['violins'] + spec['fractions']


def verdict_key(v):
    return next((k for k in VERDICT_COLOR if v.startswith(k)), 'missing')


def load(replicates):
    spec = json.loads((replicates / 'spec.json').read_text(encoding='utf-8'))
    with open(replicates / 'consistency.csv', encoding='utf-8-sig', newline='') as fh:
        rows = list(csv.DictReader(fh))
    labels = {}
    # Newest manifest wins, as volscope_compare_experiments.m labels the panels' y-axes; the
    # first experiment's label used to win here, so a relabeled metric would title its slide
    # differently from the axis beside it (bug hunt 2026-10-26).
    mans = [Path(e['comparison_dir']) / 'manifest.json' for e in spec['experiments']]
    # oldest first, newest last; reversed() first so a time tie goes to the FIRST experiment,
    # as MATLAB's max() does
    for f in sorted(reversed(mans), key=lambda f: f.stat().st_mtime):
        man = json.loads(f.read_text(encoding='utf-8'))
        for m in man['metrics'] + man.get('fraction_metrics', []):
            labels[m['name']] = m['ylabel'].replace('\u00b5', '\u03bc')   # deck mu = U+03BC
    return spec, rows, labels


def summary_slide(deck, spec, rows, labels):
    s = deck.slide('Do the two experiments agree? (CAT vs FMC, per timepoint)')
    exps = [e['name'] for e in spec['experiments']]
    tps = spec['tp_labels']
    metrics = slide_order(spec)
    by = {(r['Metric'], r['Timepoint']): r for r in rows}
    left, top, width = 0.3, 0.85, bpd.SLIDE_W - 0.6
    shape = s.shapes.add_table(len(metrics) + 1, len(tps) + 1, Inches(left), Inches(top),
                               Inches(width), Inches(0.42 * (len(metrics) + 1)))
    t = shape.table
    t.columns[0].width = Inches(4.3)
    for c in range(1, len(tps) + 1):
        t.columns[c].width = Inches((width - 4.3) / len(tps))

    def put(r, c, text, size, bold=False, fill=None):
        cell = t.cell(r, c)
        cell.text = text
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(size); run.font.bold = bold
                run.font.color.rgb = bpd.BLACK
        if fill is not None:
            cell.fill.solid(); cell.fill.fore_color.rgb = fill

    put(0, 0, 'Metric', 14, True)
    for c, tp in enumerate(tps, 1):
        put(0, c, tp, 14, True)
    for r, m in enumerate(metrics, 1):
        put(r, 0, labels.get(m, m), 12, True)
        for c, tp in enumerate(tps, 1):
            row = by.get((m, tp))
            if row is None:
                put(r, c, 'missing', 11, fill=VERDICT_COLOR['missing'])
                continue
            per = []
            for e in exps:
                x = 'e_' + e
                p = row.get(f'{x}_P_Value_Adj', '')
                per.append(f"{stars(p) if p not in ('', 'NaN') else '-'} "
                           f"{row.get(f'{x}_Direction', '')}".strip())
            verdict = row['Verdict']
            for e in exps:                          # 'sig. in 20261001' -> 'sig. in Oct 1, 2026'
                verdict = verdict.replace(e, experiment_label(e))
            put(r, c, f"{verdict}\n" + '  |  '.join(per), 11,
                fill=VERDICT_COLOR[verdict_key(row['Verdict'])])
    # Color key: two columns of two, between the table and the footnote.
    for i, (k, text) in enumerate(VERDICT_KEY):
        x = left + (i % 2) * width / 2
        y = 6.5 + (i // 2) * 0.32
        sw = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.04),
                                Inches(0.42), Inches(0.24))
        sw.fill.solid(); sw.fill.fore_color.rgb = VERDICT_COLOR[k]
        sw.line.color.rgb = RGBColor(0x80, 0x80, 0x80)
        bpd.add_text(s, x + 0.5, y, width / 2 - 0.6, 0.32, text, 14)
    bpd.add_text(s, left, bpd.SLIDE_H - 0.36, width, 0.34,
                 f"Each cell: verdict, then {' | '.join(experiment_label(e) for e in exps)} "
                 f"(BH-adjusted p, each experiment's own tests: Fisher's exact test for "
                 f"% of cells, rank-sum otherwise)", 12,
                 color=bpd.GREY, align=PP_ALIGN.CENTER)


def metric_slide(deck, spec, replicates, metric, label):
    exps = spec['experiments']
    s = deck.slide(label)
    n = len(exps)
    top, gap = 1.15, 0.2
    bw = (bpd.SLIDE_W - 0.4 - gap * (n - 1)) / n
    for i, e in enumerate(exps):
        x = 0.2 + i * (bw + gap)
        bpd.add_text(s, x, 0.66, bw, 0.45, experiment_label(e['name']), 20, bold=True,
                     align=PP_ALIGN.CENTER)
        bpd.fit_in_box(s, replicates / e['name'] / f'{metric}_by_timepoint.png',
                       x, top, bw, bpd.SLIDE_H - top - 0.1)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('--replicates', type=Path, default=REPLICATES)
    ap.add_argument('--export', action='store_true',
                    help='export slide PNGs to <replicates>/deck_png_<yyyymmdd>')
    args = ap.parse_args()
    today = datetime.date.today().strftime('%Y%m%d')
    out = DECK_DIR / f'CART_CD19_analysis_summary_{today}.pptx'
    if (out.parent / f'~${out.name}').exists():
        raise SystemExit(f'{out.name} is open in PowerPoint -- close it first')
    bpd.OUT = out                                   # Deck.save / export_pngs read OUT

    spec, rows, labels = load(args.replicates)
    deck = bpd.Deck()
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.titles.append('Title')
    bpd.add_text(s, 0.5, 2.3, bpd.SLIDE_W - 1.0, 1.0, 'CD19 internalization: fixed coculture data',
                 40, bold=True, align=PP_ALIGN.CENTER)
    bpd.add_text(s, 0.5, 3.4, bpd.SLIDE_W - 1.0, 0.6,
                 ' vs '.join(experiment_label(e['name']) for e in spec['experiments'])
                 + '  \u00b7  CAT vs FMC  \u00b7  '
                 + ' / '.join(t.replace(' min', '') for t in spec['tp_labels']) + ' min',
                 24, align=PP_ALIGN.CENTER)
    bpd.add_text(s, 0.5, 4.2, bpd.SLIDE_W - 1.0, 0.5, datetime.date.today().strftime('%B %d, %Y'),
                 20, color=bpd.GREY, align=PP_ALIGN.CENTER)

    missing = []
    for m in slide_order(spec):
        for e in spec['experiments']:
            if not (args.replicates / e['name'] / f'{m}_by_timepoint.png').is_file():
                missing.append(f"{e['name']} / {m}")
        metric_slide(deck, spec, args.replicates, m, labels.get(m, m))
    # The summary table closes the deck, after the panels it summarizes (FF 2026-10-25).
    summary_slide(deck, spec, rows, labels)
    deck.save()
    print(f'wrote {out}')
    for k in missing:
        print(f'[MISSING] {k}')
    if args.export:
        bpd.export_pngs(args.replicates / f'deck_png_{today}', deck.titles)
    return 1 if missing else 0


if __name__ == '__main__':
    sys.exit(main())
