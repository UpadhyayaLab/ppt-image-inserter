#!/usr/bin/env python3
r"""build_cd19_figure_candidates_deck.py  CD19 coculture figure candidates, for FF to pick from.

FF 2026-10-27: examples that show (1) CD19 accumulation at the contact and (2) CD19 inside
the T cell, all from the 2nd experiment (Oct 1, 2026), ONE contrast for every figure tile.
Tiles: volscope dev/cd19_engagement_figure.m (XY at the widest plane, XZ through the cell
center, XY at the plane with most internalized CD19; merge and CD19 alone, each at several
CD19 ceilings). Shortlist: volscope dev/select_cd19_figure_candidates.m (per condition, 3
'representative' cells near the median of the synapse / inside rank score, 3 'strong').

Slides: title; contrast slides (two cells, every ceiling, merge and CD19 alone) so one
ceiling is picked for the figure; then per condition the representative and the strong
cells, 3 rows x (widest merge | widest CD19 | side merge | most-puncta merge) at --cd19-hi.
Labels are slide text (bold Calibri), never in the tiles.

Run:  python build_cd19_figure_candidates_deck.py --root <engagement dir> --exp "Oct 1, 2026"
          [--cd19-hi 1500] [--export DIR]
Out:  K:\FF\PPT\PPT_autogeneration\CART\CD19\CART_CD19_figure_candidates_<today>.pptx
"""
from __future__ import annotations

import argparse
import csv
import datetime
import re
import sys
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'CART_pMLC_ppt'))
import build_pmlc_profiles_deck as bpd  # noqa: E402

DECK_DIR = Path(r'K:\FF\PPT\PPT_autogeneration\CART\CD19')
# Columns of a candidate row: (tile kind, header). 'xy_cd19only' = CD19 alone.
COLUMNS = [('xy', 'Widest plane'), ('xy_cd19only', 'Widest plane, CD19 only'),
           ('side', 'Side view (XZ)'), ('puncta', 'Most internalized CD19')]
LIST_TITLE = {'representative': 'near the median', 'strong': 'strongest'}
DASH = '—'


def cond_key(c):
    m = re.match(r'^(CAT|FMC)_(\d+)min', c)
    return (m.group(1) != 'CAT', int(m.group(2))) if m else (True, 0, c)


def cond_label(c):
    m = re.match(r'^(CAT|FMC)_(\d+)min', c)
    return f'{m.group(1)} {m.group(2)} min' if m else c


def tile(root, cond, cell, kind, hi):
    return root / cond / f'{cond}_{cell}_{kind}_cd19hi{hi}.png'


def no_objects(root, cond, cell):
    """True when the driver recorded zero internalized-object voxels for the cell."""
    with open(root / cond / f'{cond}_{cell}_info.csv', newline='') as h:
        return int(next(csv.DictReader(h))['n_object_voxels']) == 0


def ceilings(root):
    """The CD19 ceilings the tiles were written at, from any condition's tiles.csv."""
    for f in sorted(root.glob('*/*_tiles.csv')):
        with open(f, newline='') as h:
            for r in csv.DictReader(h):
                return [int(x) for x in r['cd19_ceilings'].split()]
    raise SystemExit(f'no *_tiles.csv under {root}')


def contrast_slides(deck, root, exp, cond, cell, his):
    for kind, what in (('xy', 'widest plane'), ('side', 'side view (XZ)')):
        s = deck.slide(f'Contrast: {cond_label(cond)}, {cell}, {what} {DASH} {exp}')
        ncol, top, left0 = len(his), 1.25, 1.3
        bw = (bpd.SLIDE_W - left0 - 0.2) / ncol
        bh = (bpd.SLIDE_H - top - 0.1) / 2
        for j, hi in enumerate(his):
            bpd.add_text(s, left0 + j * bw, 0.75, bw, 0.4, f'CD19 110\u2013{hi}', 16, bold=True,
                         align=PP_ALIGN.CENTER)
            for i, (sfx, _) in enumerate((('', ''), ('_cd19only', ''))):
                bpd.fit_in_box(s, tile(root, cond, cell, kind + sfx, hi),
                               left0 + j * bw, top + i * bh, bw - 0.08, bh - 0.06)
        for i, lab in enumerate(('CAR + CD19', 'CD19')):
            bpd.add_text(s, 0.1, top + i * bh + bh / 2 - 0.25, left0 - 0.15, 0.5, lab, 16,
                         bold=True, align=PP_ALIGN.RIGHT)


def candidate_slide(deck, root, exp, cond, lst, rows, hi):
    s = deck.slide(f'{cond_label(cond)}, {LIST_TITLE[lst]} cells (CD19 110\u2013{hi}) '
                   f'{DASH} {exp}')
    top, left0 = 1.2, 2.3
    bw = (bpd.SLIDE_W - left0 - 0.2) / len(COLUMNS)
    bh = (bpd.SLIDE_H - top - 0.1) / 3
    for j, (_, head) in enumerate(COLUMNS):
        bpd.add_text(s, left0 + j * bw, 0.72, bw, 0.45, head, 14, bold=True, align=PP_ALIGN.CENTER)
    missing = []
    for i, r in enumerate(rows[:3]):
        y = top + i * bh
        bpd.add_text(s, 0.1, y + 0.15, left0 - 0.2, bh - 0.2,
                     [r['cell'], f"synapse CD19: p{r['rank_synapse']}",
                      f"CD19 inside: p{r['rank_inside']}", f"of {r['n_ranked']} cells"],
                     14, bold=True, align=PP_ALIGN.RIGHT)
        for j, (kind, _) in enumerate(COLUMNS):
            p = tile(root, cond, r['cell'], kind, hi)
            if kind == 'puncta' and not p.is_file() and no_objects(root, cond, r['cell']):
                # by design: the driver writes no puncta tile for a cell with no kept object
                bpd.add_text(s, left0 + j * bw, y + bh / 2 - 0.25, bw - 0.08, 0.5,
                             'no internalized CD19', 14, bold=True, color=bpd.GREY,
                             align=PP_ALIGN.CENTER)
                continue
            if not p.is_file():
                missing.append(p.name)
            bpd.fit_in_box(s, p, left0 + j * bw, y, bw - 0.08, bh - 0.06)
    return missing


CHANNELS = [('_caronly', 'CAR'), ('_cd19only', 'CD19'), ('', 'CAR + CD19')]
# candidates.csv list -> slide wording (FF 2026-10-28, his words: "T cell on top of target?
# next to it?"); the class limits live in the run script (synapse_normal_angle_deg)
RULE_CELLS = ('Cells: target visible in both views (≥ 10 µm² of CD19 ≥ 2 µm from the '
              'T cell), FMC projected area ≥ CAT’s, then closest to the condition median '
              '(FMC 60 min next to target: C2-cell12, chosen by eye).')
PHENOTYPE = {'on_top': 'T cell on top of target', 'next_to': 'T cell next to target'}
# Angle class limits (synapse_normal_angle_deg, deg from straight down), the `classes` of
# TCell dev/runs/run_cd19_fig_allcells_20261028.m. Both the check of each --pair-cells
# phenotype and the slide's rule line read these, so the text cannot drift from the check.
PHENO_DEG = {'on_top': (0, 45), 'next_to': (70, 110)}
VIEWS = [('xy', 'Widest plane'), ('side', 'Side view (XZ)')]


def read_info(root, cond, cell):
    with open(root / cond / f'{cond}_{cell}_info.csv', newline='') as h:
        return next(csv.DictReader(h))


def car_windows(info):
    """'110–187' when both views share a window, else 'XY 110–187 · XZ 110–186'."""
    xy = info['car_window'].replace('-', '–')
    side = (info.get('car_window_side') or info['car_window']).replace('-', '–')
    return xy if side == xy else f'XY {xy} · XZ {side}'


def car_rule_text(rule):
    """The CAR sentence of a slide, from the rule cd19_engagement_figure.m recorded."""
    m = re.match(r'^(view_|mask_)?p([\d.]+)(_nosyn)?$', rule)
    if rule == 'fixed':
        return 'CAR: one fixed window for every cell.'
    if not m:
        raise SystemExit(f'unknown CAR rule {rule!r}: add its wording to car_rule_text')
    kind, pct, nosyn = m.groups()
    if kind == 'view_':
        return f'CAR: 110 to the {pct}th percentile inside the T cell outline, set per view.'
    if kind == 'mask_':
        return f'CAR: 110 to the {pct}th percentile inside the T cell, one window per cell.'
    return (f'CAR: 110 to the cell’s own {pct}th percentile'
            + (', synapse left out.' if nosyn else '.'))


def pair_angle(root, cond, cell):
    """synapse_normal_angle_deg of a pair cell, from the root's synapse_angle.csv."""
    f = root / 'synapse_angle.csv'
    if not f.is_file():
        raise SystemExit(f'{f} missing: the on top / next to label cannot be checked')
    with open(f, newline='') as h:
        for r in csv.DictReader(h):
            if r['condition'] == cond and r['cell'] == cell:
                return float(r['synapse_normal_angle_deg'])
    raise SystemExit(f'{cond} {cell} not in {f}')


def channel_slide(deck, root, exp, r, hi):
    """FF 2026-10-28: "both channels then merge with both views" -- one cell per slide,
    columns CAR | CD19 | merge, rows widest plane | side view."""
    cond, cell = r['condition'], r['cell']
    pheno = PHENOTYPE.get(r['list'], '')
    s = deck.slide(f'{cond_label(cond)}, {cell}{", " + pheno if pheno else ""} '
                   f'(CD19 110–{hi}) {DASH} {exp}')
    top, left0 = 1.2, 2.6
    bh = (bpd.SLIDE_H - top - 0.1) / len(VIEWS)
    bw = min((bpd.SLIDE_W - left0 - 0.2) / len(CHANNELS), bh)
    for j, (_, head) in enumerate(CHANNELS):
        bpd.add_text(s, left0 + j * bw, 0.72, bw, 0.45, head, 16, bold=True, align=PP_ALIGN.CENTER)
    missing = []
    for i, (view, vlab) in enumerate(VIEWS):
        y = top + i * bh
        bpd.add_text(s, 0.1, y + bh / 2 - 0.25, left0 - 0.2, 0.5, vlab, 16, bold=True,
                     align=PP_ALIGN.RIGHT)
        for j, (sfx, _) in enumerate(CHANNELS):
            p = tile(root, cond, cell, view + sfx, hi)
            if not p.is_file():
                missing.append(p.name)
                continue
            bpd.fit_in_box(s, p, left0 + j * bw, y, bw - 0.08, bh - 0.06)
    # both views' windows: the side tiles use car_window_side, which can differ from XY's
    car = car_windows(read_info(root, cond, cell))
    bpd.add_text(s, 0.1, top + bh / 2 + 0.4, left0 - 0.2, 1.6,
                 [f"synapse CD19: p{r['rank_synapse']}", f"CD19 inside: p{r['rank_inside']}",
                  f"of {r['n_ranked']} cells", f'CAR {car}']
                 + ([f"synapse {r['direction_deg']}° from down"]
                    if r.get('direction_deg') not in (None, '', 'NaN') else []), 14, bold=True,
                 color=bpd.GREY, align=PP_ALIGN.RIGHT)
    return missing


# Channel names in their own display colors (FF 2026-10-28: "CAR label should be cyan CD19
# magenta"), the config's #40e0ff / #ff4fd8, the tiles' tints (cd19_engagement_figure.m).
CHANNEL_RGB = {'CAR': RGBColor(0x40, 0xE0, 0xFF), 'CD19': RGBColor(0xFF, 0x4F, 0xD8)}


def channel_header(slide, left, top, width, height, text, size):
    """A centered bold header with each channel name ('CAR', 'CD19') in its color."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    para = tb.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    for part in re.split(r'(CAR|CD19)', text):
        if not part:
            continue
        r = para.add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = CHANNEL_RGB.get(part, bpd.BLACK)


def pair_slides(args, rows):
    """--layout pairs (FF 2026-10-28: "one slide per timepoint, CAT FMC comparison for each.
    1 rep cell each view"): --pair-cells cond:cell,... ; per time point CAT | FMC side by side,
    each CAR | CD19 | merge, rows = the two views; the CD19 window per --cd19-hi-time."""
    by_time = dict(kv.split('=') for kv in args.cd19_hi_time.split(',') if kv)
    # cond:cell[:phenotype]; a phenotype (on_top / next_to) gives one slide per (time, phenotype)
    want = [(p.split(':') + [''])[:3] for p in args.pair_cells.split(',') if p]
    drop = {tuple(d.split(':', 1)) for d in args.drop.split(',') if d}
    want = [w for w in want if (w[0], w[1]) not in drop]   # --drop holds here too
    for c, l, ph in want:
        if ph and ph not in PHENO_DEG:
            raise SystemExit(f'{c}:{l}: phenotype {ph!r} is not one of {sorted(PHENO_DEG)}')
        if ph:                           # the typed label must match the measured angle
            a, (lo, hi_deg) = pair_angle(args.root, c, l), PHENO_DEG[ph]
            if not lo <= a <= hi_deg:
                raise SystemExit(f'{c} {l}: synapse {a:.0f}° from down, outside {ph} '
                                 f'({lo}–{hi_deg}°)')
    rules_car = {read_info(args.root, c, l).get('car_rule') or '' for c, l, _ in want}
    if rules_car == {''}:                # tiles from before car_rule was recorded
        if not args.car_rule:
            raise SystemExit('these tiles do not record their CAR rule: pass --car-rule '
                             '(the run script\'s opts.car_win, e.g. view_p95)')
        rules_car = {args.car_rule}
    elif '' in rules_car or len(rules_car) > 1 or (args.car_rule and {args.car_rule} != rules_car):
        raise SystemExit(f'CAR rules disagree: tiles {sorted(rules_car)}, --car-rule {args.car_rule}')
    car_sentence = car_rule_text(rules_car.pop())
    groups = sorted({(cond_key(c)[1], ph) for c, _, ph in want},
                    key=lambda g: (g[0], g[1] != 'on_top'))
    deck = bpd.Deck()
    missing = []
    for t, ph in groups:
        hi = int(by_time.get(str(t), args.cd19_hi))
        cells = sorted([(c, l) for c, l, p in want if cond_key(c)[1] == t and p == ph],
                       key=lambda w: cond_key(w[0]))
        if [c.split('_')[0] for c, _ in cells] != ['CAT', 'FMC']:
            raise SystemExit(f'{t} min {ph or ""}: needs one CAT and one FMC cell, got '
                             f'{cells} (--drop both to leave the slide out)')
        what = f', {PHENOTYPE[ph]}' if ph in PHENOTYPE else ''
        # no experiment in the title; the window worded as the montage deck's (FF 2026-10-28)
        s = deck.slide(f'{t} min{what}: CAT vs FMC (Fixed Contrast CD19 110–{hi})')
        top, left0, ggap = 1.6, 2.15, 0.3
        bw = min((bpd.SLIDE_W - left0 - 0.15 - ggap) / (len(CHANNELS) * len(cells)),
                 (bpd.SLIDE_H - top - 0.5) / len(VIEWS))
        bh = bw + 0.04                       # rows packed: square tiles, a small gap
        for g, (cond, cell) in enumerate(cells):
            x0 = left0 + g * (len(CHANNELS) * bw + ggap)
            bpd.add_text(s, x0, 0.68, len(CHANNELS) * bw, 0.42,
                         f"{cond_label(cond).split()[0]} {DASH} {cell}", 18, bold=True,
                         align=PP_ALIGN.CENTER)
            for j, (sfx, head) in enumerate(CHANNELS):
                channel_header(s, x0 + j * bw, 1.1, bw, 0.4, head, 14)
                for i, (view, _) in enumerate(VIEWS):
                    p = tile(args.root, cond, cell, view + sfx, hi)
                    if not p.is_file():
                        missing.append(p.name)
                        continue
                    bpd.fit_in_box(s, p, x0 + j * bw, top + i * bh, bw - 0.06, bh - 0.06)
            # (no per-cell CAR window / angle line under the tiles, FF 2026-10-28: "can be
            # removed"; the values stay in each cell's _info.csv and synapse_angle.csv)
        for i, (_, vlab) in enumerate(VIEWS):
            bpd.add_text(s, 0.05, top + i * bh + bh / 2 - 0.3, left0 - 0.12, 0.6, vlab, 14,
                         bold=True, align=PP_ALIGN.RIGHT)
        # The rules, on the slide (FF 2026-10-28: "these rules should be on the slides
        # (written concisely)"); the empty band under the tiles.
        rules = [f'CD19: 110–{hi}, the same for CAT and FMC. {car_sentence}',
                 ('XY: the plane with the largest synapse area. XZ: the slice with the most '
                  'synapse.')]
        if ph in PHENOTYPE:
            rules.append(f"On top of target: synapse within {PHENO_DEG['on_top'][1]}° of "
                         f"straight down. Next to target: {PHENO_DEG['next_to'][0]}–"
                         f"{PHENO_DEG['next_to'][1]}°.")
        if args.rule_cells:                 # how the cells were picked (run script's rule)
            rules.append(RULE_CELLS)
        bpd.add_text(s, 0.4, top + len(VIEWS) * bh + 0.55, bpd.SLIDE_W - 0.8, 1.0, rules, 13,
                     bold=True, color=bpd.GREY, align=PP_ALIGN.LEFT)
    return deck, missing


def main_channels(args, rows):
    """--layout channels: a title, then one channel_slide per candidate, conditions in order."""
    deck = bpd.Deck()
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.titles.append('Title')
    bpd.add_text(s, 0.5, 2.3, bpd.SLIDE_W - 1.0, 1.0, 'CD19 figure candidates', 40, bold=True,
                 align=PP_ALIGN.CENTER)
    bpd.add_text(s, 0.5, 3.4, bpd.SLIDE_W - 1.0, 0.6,
                 f'{args.exp} · CD19 window on each slide (one per time point) · '
                 'CAR per cell', 24,
                 align=PP_ALIGN.CENTER)
    by_time = dict(kv.split('=') for kv in args.cd19_hi_time.split(',') if kv)
    missing = []
    for r in sorted(rows, key=lambda r: (cond_key(r['condition']), r['list'] != 'forced')):
        t = str(cond_key(r['condition'])[1])
        missing += channel_slide(deck, args.root, args.exp, r, int(by_time.get(t, args.cd19_hi)))
    return deck, missing


def window_slides(args, rows):
    """--layout windows (FF 2026-10-28, windows anchored on cell 12): <root>/windows.csv
    (pct, car, cd19) and tiles under <root>/p<pct>/<cond>/. One slide per shown cell (the
    forced cell, then each condition's first 'trend' cell), view = --view; columns are the
    windows, rows CAR | CD19 | merge."""
    with open(args.root / 'windows.csv', newline='') as h:
        wins = list(csv.DictReader(h))
    shown, seen = [], set()
    for r in sorted(rows, key=lambda r: (r['list'] != 'forced', cond_key(r['condition']))):
        if r['condition'] not in seen:
            shown.append(r)
            seen.add(r['condition'])
    shown.sort(key=lambda r: cond_key(r['condition']))
    deck = bpd.Deck()
    missing = []
    vlab = dict(VIEWS)[args.view]
    for r in shown:
        cond, cell = r['condition'], r['cell']
        s = deck.slide(f'{cond_label(cond)}, {cell}, {vlab.lower()} {DASH} {args.exp}')
        top, left0 = 1.25, 1.6
        bh = (bpd.SLIDE_H - top - 0.1) / len(CHANNELS)
        bw = min((bpd.SLIDE_W - left0 - 0.2) / len(wins), bh * 1.0)
        for j, w in enumerate(wins):
            bpd.add_text(s, left0 + j * bw, 0.62, bw, 0.6,
                         [f"cell 12 p{w['pct']}",
                          f"CAR 110–{w['car']} · CD19 110–{w['cd19']}"], 13, bold=True, align=PP_ALIGN.CENTER)
            for i, (sfx, _) in enumerate(CHANNELS):
                p = (args.root / f"p{w['pct']}" / cond /
                     f"{cond}_{cell}_{args.view}{sfx}_cd19hi{w['cd19']}.png")
                if not p.is_file():
                    missing.append(p.name)
                    continue
                bpd.fit_in_box(s, p, left0 + j * bw, top + i * bh, bw - 0.08, bh - 0.06)
        for i, (_, lab) in enumerate(CHANNELS):
            bpd.add_text(s, 0.1, top + i * bh + bh / 2 - 0.25, left0 - 0.15, 0.5, lab, 16,
                         bold=True, align=PP_ALIGN.RIGHT)
    return deck, missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--root', required=True, type=Path,
                    help='<experiment>/cd19_figures/engagement_<date>, one folder per condition')
    ap.add_argument('--candidates', type=Path, default=None,
                    help='candidates.csv (default <root>/candidates.csv)')
    ap.add_argument('--exp', required=True, help='experiment label, e.g. "Oct 1, 2026"')
    ap.add_argument('--cd19-hi', type=int, default=1500, help='CD19 ceiling of the candidate slides')
    ap.add_argument('--contrast-cells', nargs='*', default=['FMC_60min:C2-cell12'],
                    help='cond:cell for the contrast slides; the strongest CAT_60min cell is added')
    ap.add_argument('--export', type=Path, default=None, help='also export slide PNGs here')
    ap.add_argument('--pair-cells', default='',
                    help='--layout pairs: cond:cell,... one CAT and one FMC cell per time point')
    ap.add_argument('--layout', choices=['grid', 'channels', 'windows', 'pairs'], default='grid',
                    help='grid: contrast + 3 cells per slide; channels: one cell per slide, '
                         'CAR | CD19 | merge x widest | side; windows: one cell per condition '
                         'across the windows of <root>/windows.csv')
    ap.add_argument('--view', choices=['xy', 'side'], default='xy',
                    help='--layout windows: which view')
    ap.add_argument('--drop', default='',
                    help='cond:cell,... left out of the deck (FF 2026-10-28: drop cells where '
                         '"we see very little target")')
    ap.add_argument('--xy-label', default=None,
                    help='row label of the XY tiles (default "Widest plane"); the renderer\'s '
                         'xy_plane decides which plane they are')
    ap.add_argument('--rule-cells', action='store_true',
                    help='--layout pairs: add the RULE_CELLS footer line (the target-visible, '
                         'closest-to-median pick of run_cd19_fig_allcells_20261028.m)')
    ap.add_argument('--car-rule', default='',
                    help='--layout pairs: the CAR rule (opts.car_win, e.g. view_p95) for tiles '
                         'whose _info.csv predates its car_rule column; checked against it '
                         'when present')
    ap.add_argument('--side-label', default=None,
                    help='row label of the XZ tiles (default "Side view (XZ)")')
    ap.add_argument('--cd19-hi-time', default='',
                    help='--layout channels: CD19 ceiling per time point, e.g. "10=300" '
                         '(FF 2026-10-28: one window per time point, CAT and FMC the same)')
    ap.add_argument('--out-name', default='CART_CD19_figure_candidates',
                    help='deck file name stem; _<today>.pptx is appended')
    args = ap.parse_args()
    if args.xy_label:
        VIEWS[0] = ('xy', args.xy_label)
    if args.side_label:
        VIEWS[1] = ('side', args.side_label)
    cand = args.candidates or args.root / 'candidates.csv'
    with open(cand, newline='') as h:
        rows = list(csv.DictReader(h))
    drop = {tuple(d.split(':', 1)) for d in args.drop.split(',') if d}
    rows = [r for r in rows if (r['condition'], r['cell']) not in drop]
    today = datetime.date.today().strftime('%Y%m%d')
    bpd.OUT = DECK_DIR / f'{args.out_name}_{today}.pptx'
    if (bpd.OUT.parent / f'~${bpd.OUT.name}').exists():
        raise SystemExit(f'{bpd.OUT.name} is open in PowerPoint -- close it first')
    if args.layout != 'windows':
        his = ceilings(args.root)
        if args.cd19_hi not in his:
            raise SystemExit(f'--cd19-hi {args.cd19_hi} not rendered; the tiles have {his}')
    if args.layout in ('channels', 'windows', 'pairs'):
        deck, missing = {'channels': main_channels, 'windows': window_slides,
                         'pairs': pair_slides}[args.layout](args, rows)
        deck.save()
        print(f'saved {bpd.OUT} ({len(deck.prs.slides)} slides)')
        for m in missing:
            print(f'  missing: {m}')
        if args.export:
            bpd.export_pngs(args.export, deck.titles)
        return 1 if missing else 0

    deck = bpd.Deck()
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.titles.append('Title')
    bpd.add_text(s, 0.5, 2.3, bpd.SLIDE_W - 1.0, 1.0, 'CD19 figure candidates', 40, bold=True,
                 align=PP_ALIGN.CENTER)
    bpd.add_text(s, 0.5, 3.4, bpd.SLIDE_W - 1.0, 0.6,
                 f'{args.exp} \u00b7 CAT vs FMC \u00b7 CD19 at the contact and inside the T cell',
                 24, align=PP_ALIGN.CENTER)
    bpd.add_text(s, 0.5, 4.2, bpd.SLIDE_W - 1.0, 0.5,
                 datetime.date.today().strftime('%B %d, %Y'), 20, color=bpd.GREY,
                 align=PP_ALIGN.CENTER)

    pairs = [c.split(':', 1) for c in args.contrast_cells]
    cat60 = [r for r in rows if r['condition'] == 'CAT_60min' and r['list'] == 'strong']
    if cat60:
        pairs.append(['CAT_60min', cat60[0]['cell']])
    for cond, cell in pairs:
        contrast_slides(deck, args.root, args.exp, cond, cell, his)

    missing = []
    for cond in sorted({r['condition'] for r in rows}, key=cond_key):
        for lst in ('representative', 'strong'):
            sel = [r for r in rows if r['condition'] == cond and r['list'] == lst]
            if sel:
                missing += candidate_slide(deck, args.root, args.exp, cond, lst, sel, args.cd19_hi)
    deck.save()
    print(f'saved {bpd.OUT} ({len(deck.prs.slides)} slides)')
    for m in missing:
        print(f'  missing: {m}')
    if args.export:
        bpd.export_pngs(args.export, deck.titles)
    return 1 if missing else 0


if __name__ == '__main__':
    sys.exit(main())
