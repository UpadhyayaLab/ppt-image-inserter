#!/usr/bin/env python3
r"""build_grant_cd19_slide.py  One-slide grant-report deck (FF 2026-09-29): CD19
internalization by CAR T cells, CAT vs FMC -- a SEPARATE deck in the exact format of the
pMLC slide of CART_grant_actin_pMLC_<date>.pptx (build_grant_deck.py, whose helpers are
imported here and whose deck is never touched).

  A reps (CAR | CD19 | Merge), CAT C2-cell70 / FMC C2-cell30, 60 min,
    one field of view for both cells (same physical scale), per-cell CD19 range
    <- volscope dev/show_internalized_cells.m, layout 'panels'
  B % of cells with >= 1 internalized CD19 punctum, 10/30/60 min
    <- volscope comparison (MATLAB plot_fraction_bars), placed as rendered

Kiet's Mar 10, 2026 coculture (one experiment, nothing pooled).

Output: K:\FF\PPT\PPT_autogeneration\CART\grant\CART_grant_CD19_internalization_<yyyymmdd>.pptx
Run:    python build_grant_cd19_slide.py [--export <dir>]
"""
from __future__ import annotations

import argparse
import csv
import datetime
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_grant_deck as bgd  # noqa: E402  (helpers only; its main() is not run)
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

bpd = bgd.bpd
EXP = Path(r'\\10.225.56.176\ULab_data\User_data\Kiet\03102026_coculture_hek_CD19gfp_Hoescht_CAR_HA_APC')
REPS = Path(r'K:\FF\volscope_show\kiet_FMC_internalized\grant_panels_20260929_fmc30')
# restyled copy of comparison/fractions/cd19_patchfrac_ge1.fig (volscope dev/restyle_fraction_fig.m)
FRAC = REPS / 'cd19_patchfrac_ge1_slide.png'
FRAC_LABELS = REPS / 'cd19_patchfrac_ge1_slide_labels.csv'   # the bar labels MATLAB drew
OUT = bgd.OUT_DIR / f'CART_grant_CD19_internalization_{datetime.date.today():%Y%m%d}.pptx'

CYAN = RGBColor(0x00, 0xA8, 0xC8)       # CAR header; the image cyan (#40e0ff) is unreadable on white
X0 = bgd.X0


def frac_numbers(s, csv_path, x, y, w, h, size=16):
    """One line per timepoint -- '10 min:  CAT 51% (25/49)   FMC 74% (78/105)' -- each cell
    type in its bar color. Strings are copied verbatim from the labels MATLAB drew on the
    bars (restyle_fraction_fig.m's _labels.csv); nothing is recomputed here."""
    rows = list(csv.DictReader(open(csv_path, encoding='utf-8')))
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for mins in sorted({int(float(r['minutes'])) for r in rows}):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = PP_ALIGN.LEFT        # FF 2026-09-29: timepoints start at one x
        runs = [(f'{mins} min:  ', bpd.BLACK)]
        for r in (r for r in rows if int(float(r['minutes'])) == mins):
            runs.append((f"{r['cell_type']} {r['pct']} ({r['n']})   ", RGBColor.from_string(r['color'])))
        for text, col in runs:
            run = para.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = col


def cd19_slide(deck):
    # FF 2026-09-29: no panel letters
    s = deck.slide('CD19 Internalization by CAR T Cells')     # FF 2026-09-29: no "60 min"
    tile = 1.62
    # FF 2026-09-29: no Internalized column
    bgd.reps_block(s, REPS, '{c}_60min*_{suffix}.png',
                   [('car', 'CAR', CYAN), ('cd19', 'CD19', bgd.MAGENTA),
                    ('merge', 'Merge', bpd.BLACK)],
                   X0, 0.68, 1.05, tile)
    # B: % of cells with >= 1 internalized punctum, right after A's three columns
    xb = X0 + 0.95 + 3 * (tile + 0.06) + 0.45
    bpd.fit_in_box(s, FRAC, xb, 0.8, 12.95 - xb, 3.55)
    # B's per-bar numbers as text under the plot (FF 2026-09-29: larger, off the bars)
    frac_numbers(s, FRAC_LABELS, xb + 0.35, 4.45, 12.95 - xb - 0.35, 1.1)
    # comparison/statistics_summary.csv (cd19_patchfrac_ge1, BH-adjusted Fisher)
    bgd.bullets(s, X0, 5.45, 12.4, 1.4,          # FF 2026-09-29: a bit lower
                ['\u2022 One experiment so far: both CAR T cell types internalize target '
                 'CD19 into deep puncta',
                 '\u2022 Most cells contain \u22651 internalized punctum by 60 min: '
                 'CAT 91%, FMC 93%',
                 '\u2022 Internalization appears to be higher in FMC cells at the '
                 '10-minute time point (74% vs. 51%)',
                 '• n = 1 experiment; cells shown are at the 60-minute time point'])
    bgd.footer(s, 'One experiment (Mar 10, 2026; n = 49/95/78 CAT and 105/109/86 FMC cells at '
                  '10/30/60 min). Internalized punctum: CD19 object \u22651 \u03bcm inside the '
                  'T cell, not connected to its membrane. Representative cells at 60 min, '
                  'single confocal slices, each contrast-scaled individually.',
               y=6.85, h=0.6)
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--export', help='also export one PNG per slide to this dir')
    a = ap.parse_args()
    if OUT.exists() and bgd.backup_presentation is not None:
        bgd.backup_presentation(str(OUT), backup_base=str(bgd.OUT_DIR / 'backups'))
    deck = bpd.Deck()
    cd19_slide(deck)
    bgd.calibri(deck.prs)
    deck.prs.save(str(OUT))
    print(f'saved {OUT} ({len(deck.prs.slides)} slides)')
    if a.export:
        bpd.OUT = OUT                       # export_pngs opens bpd.OUT
        bpd.export_pngs(a.export, deck.titles)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
