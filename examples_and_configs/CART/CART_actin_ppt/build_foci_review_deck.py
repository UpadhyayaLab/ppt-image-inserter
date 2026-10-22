#!/usr/bin/env python3
r"""build_foci_review_deck.py  Separate review deck of every actin foci metric (+ the clearance
area it is normalized by) for the 2026-07 confocal actin-only CAR-T data (Kiet; Jul 17 + Jul 23;
CAT / FMC x 3/5/7/10 min), FF 2026-09-17: "construct separate PPT which I will review".

Includes the metrics added Aug 25-27 (Dalby et al. 2026): actin_foci_density_um2,
actin_foci_density_clearance_um2, actin_foci_median_area_um2. Nothing is plotted here; panels
come from the Sep 15 compile + per-cell montages (first montages_filtered page).

Output: K:\FF\PPT\PPT_autogeneration\CART\actin_only\CART_actin_foci_review_20260717_20260723_<yyyymmdd>.pptx
Run:    python build_foci_review_deck.py [--export <dir>]
PowerPoint is only touched by --export (never quit; skipped if FF has the deck open).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import build_actin_decks as bad
from build_actin_decks import bpd

TS = bad.TS
MU2, PER_MU2 = '\u03bcm\u00b2', '\u03bcm\u207b\u00b2'
# (title, grid panel relative to <compile>\grid_panels\, one-line definition for the overview)
CLEARANCE = [
    ('Actin Clearance Area', f'{TS}/actin_bottom_slice_clearance_area_um2_grid.png',
     f'Largest connected low-actin region inside the synapse mask (smoothed actin < 0.25 x mask mean; holes filled), {MU2}'),
    ('Actin Clearance Area Fraction', f'{TS}/actin_bottom_slice_clearance_frac_grid.png',
     'Clearance area / synapse mask area'),
]
FOCI = [
    ('Actin Foci Count', f'{TS}/actin_foci_count_grid.png', 'Number of connected foci in the synapse mask'),
    ('Actin Foci Area', f'{TS}/actin_foci_area_um2_grid.png', f'Summed area of all foci, {MU2}'),
    ('Actin Foci Area Fraction', f'{TS}/actin_foci_area_fraction_grid.png', 'Summed foci area / synapse mask area'),
    ('Actin Foci Density (NEW)', f'{TS}/actin_foci_density_um2_grid.png',
     f'Foci count / synapse mask area, {PER_MU2}'),
    ('Actin Foci per Clearance Area (NEW)', f'{TS}/actin_foci_density_clearance_um2_grid.png',
     f'Foci count / clearance area, {PER_MU2}  (Dalby convention; NaN when no clearance)'),
    ('Actin Foci Median Area (NEW)', f'{TS}/actin_foci_median_area_um2_grid.png',
     f'Median area of a single focus, {MU2}'),
]
INTENSITY = [
    ('Actin Foci Mean Intensity', f'{TS}/actin_foci_mean_intensity_grid.png', 'Mean raw actin intensity over foci pixels'),
    ('Actin Foci Total Intensity', f'{TS}/actin_foci_total_intensity_grid.png', 'Summed raw actin intensity over foci pixels'),
    ('Actin Foci Mean Normalized Score', f'{TS}/actin_foci_mean_norm_score_grid.png',
     'Mean local-density-normalized band-pass score over foci pixels'),
    ('Actin Foci Max Normalized Score', f'{TS}/actin_foci_max_norm_score_grid.png',
     'Max local-density-normalized band-pass score over foci pixels'),
]
MONTAGES = [('Actin Foci Overlay', r'foci\final_overlay'), ('Actin Clearance', r'actin\synapse\clearance')]


def overview(deck):
    s = deck.slide(f'Actin Foci Metrics {bad.DASH} Definitions')
    rows = [('Metric', 'Definition')] + [(t, d) for t, _, d in CLEARANCE + FOCI + INTENSITY]
    tbl = s.shapes.add_table(len(rows), 2, Inches(0.4), Inches(0.8), Inches(bpd.SLIDE_W - 0.8),
                             Inches(0.46 * len(rows))).table
    tbl.columns[0].width = Inches(3.9)
    tbl.columns[1].width = Inches(bpd.SLIDE_W - 0.8 - 3.9)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = v
            for para in tbl.cell(i, j).text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(13)
                    run.font.bold = i == 0 or j == 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--compile', help='compile dir (default: newest compiled_20260717_20260723_2*)')
    ap.add_argument('--export', help='also export one PNG per slide (named by title) under this dir')
    args = ap.parse_args()
    compile_dir = Path(args.compile) if args.compile else bad.newest('compiled_20260717_20260723_2*')
    if compile_dir is None:
        print('[error] no compile found'); return 1
    bpd.OUT = bad.OUT_DIR / f'CART_actin_foci_review_20260717_20260723_{bad.TODAY}.pptx'
    lock = bpd.OUT.with_name('~$' + bpd.OUT.name)
    if lock.exists():
        print(f'[error] {bpd.OUT.name} is open in PowerPoint; close it first'); return 1
    print(f'compile: {compile_dir}')

    deck = bpd.Deck()
    overview(deck)
    missing = []
    for title, rel, _ in CLEARANCE + FOCI + INTENSITY:
        png = compile_dir / 'grid_panels' / rel
        if not png.is_file():
            missing.append(png)
        bad.single(deck, title, png)

    # Montages at ONE physical scale per view (bad.montage_group_ppi): find them all first.
    pngs = {(stem, folder, ct, t): bad.first_page(bad.NF / folder / f'{ct}_{t}' / 'split_channels'
                                                  / 'prog_fixed_cells' / sub)
            for stem, sub in MONTAGES for _, _, _, folder in bad.DATASETS for t in bad.TPS
            for ct in ('CAT', 'FMC')}
    missing += [f'{folder} {ct}_{t} {stem}' for (stem, folder, ct, t), p in pngs.items() if p is None]
    ppi = bad.montage_group_ppi({stem: [p for k, p in pngs.items() if k[0] == stem] for stem, _ in MONTAGES})
    for stem, _ in MONTAGES:
        for _, _, date, folder in bad.DATASETS:
            for t, tl in zip(bad.TPS, bad.TP_LBL):
                s = deck.slide(f'{stem} {bad.DASH} {date}, {tl}')
                for k, (ct, col) in enumerate((('CAT', bpd.CAT_BLUE), ('FMC', bpd.FMC_RED))):
                    left = 0.1 + k * (bad.MONT_HALF + 0.1)
                    bpd.add_text(s, left, 0.62, bad.MONT_HALF, 0.4, ct, 20, bold=True, color=col,
                                 align=PP_ALIGN.CENTER)
                    bad.place_montage(s, pngs[(stem, folder, ct, t)], left, ppi.get(stem))
    deck.save()
    print(f'saved {bpd.OUT} ({len(deck.prs.slides)} slides)')
    for m in missing:
        print(f'[missing] {m}')
    if args.export:
        bpd.export_pngs(Path(args.export), deck.titles)
    return 1 if missing else 0


if __name__ == '__main__':
    sys.exit(main())
