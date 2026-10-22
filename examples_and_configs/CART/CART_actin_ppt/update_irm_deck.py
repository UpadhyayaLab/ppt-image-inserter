#!/usr/bin/env python3
r"""update_irm_deck.py  Swap every plot + montage picture in FF's hand-built
CART_actin_IRM_3_5_7_10_20260916.pptx for the current data, IN PLACE, into a NEW copy.

FF 2026-09-16: the deck's pictures came from the pre-removal Aug 11 run
(Y:\Data_to_share\FF\results_compiled\actin_only\compiled_20260717_20260723_20260811 and the old H:
per-cell montages). Kiet has since removed 721 cells and the data were reprocessed (Sep 2) and
recompiled (compiled_20260717_20260723_20260915). Rules from FF:
  * new copy (source stays byte-identical); NO text edits;
  * each picture keeps its slot: aspect-fit inside its old bounding box, centred, same z-order;
  * new actin profile slides inserted after slide 48 (profile+gradient per experiment, time course,
    normalized + raw zone MFI; NOT FMC-CAT, NOT synapse MFI).
Montage slides -> page 1 of the matching prog_fixed_cells family (build_actin_decks.MONTAGES);
plot slides -> same-named panel in the new compile.
Then verifies: text/shape XML of every original slide unchanged, every picture blob == its intended
source file and != the old blob, slide count, source untouched. PowerPoint is never opened.

Run: python update_irm_deck.py
"""
from __future__ import annotations

import copy
import hashlib
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_actin_decks as bad  # noqa: E402
from build_actin_decks import bpd  # noqa: E402

SRC = Path(r'K:\FF\PPT\PPT_autogeneration\CART\actin_only\CART_actin_IRM_3_5_7_10_20260916.pptx')
TWIN = Path(r'K:\FF\PPT\CAR_T_PPT\CART_actin_IRM_3_5_7_10_20260916.pptx')
OUT = SRC.with_name(SRC.stem + '_newcompile.pptx')
COMPILE = bad.RC / 'compiled_20260717_20260723_20260915'
PANELS = bad.RC / 'actin_extended_panels_20260916'
INSERT_AFTER = 48                 # 1-based slide number the new profile slides follow
EMU_IN = 914400
TITLE_RE = re.compile(r'^(?P<stem>.+?)\s+[\u2014\u2013-]\s+(?P<date>\d{8}), (?P<tp>\d+) min\s*$')
DATE_FOLDER = {date: folder for _, _, date, folder in bad.DATASETS}


def sha(b):
    return hashlib.sha1(b).hexdigest()


def descr_of(pic):
    return pic._element.xpath('./p:nvPicPr/p:cNvPr')[0]


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ''


def plot_source(name):
    if name.startswith('axial_overlay_perslice_'):
        return COMPILE / 'axial_profile_overlays' / name
    if '_rad_profile_' in name:
        return COMPILE / 'grid_panels' / 'CAT_vs_FMC_by_timepoint' / name
    return COMPILE / 'grid_panels' / 'timecourse_scatter_plots' / name


def resolve(prs):
    """(slide_idx0, pic_shape, new_path) for every picture; abort on anything unresolved."""
    fam = {stem.lower(): sub for stem, sub in bad.MONTAGES}
    plan, errors = [], []
    for i, s in enumerate(prs.slides):
        pics = sorted((sh for sh in s.shapes if sh.shape_type == 13), key=lambda p: p.left)
        if not pics:
            continue
        m = TITLE_RE.match(title_of(s))
        if m:                                             # montage slide: left = CAT, right = FMC
            sub = fam.get(m['stem'].lower())
            if sub is None or len(pics) != 2:
                errors.append(f'slide {i + 1}: montage family/picture count unresolved ({m["stem"]!r}, {len(pics)})')
                continue
            for ct, pic in zip(('CAT', 'FMC'), pics):
                centre = (pic.left + pic.width / 2) / EMU_IN
                if (ct == 'CAT') != (centre < bpd.SLIDE_W / 2):
                    errors.append(f'slide {i + 1}: {ct} picture on the wrong half')
                folder = (bad.NF / DATE_FOLDER[m['date']] / f'{ct}_{m["tp"]}min' / 'split_channels'
                          / 'prog_fixed_cells' / sub)
                new = bad.first_page(folder)
                old = descr_of(pic).get('descr') or ''
                if new is None:
                    errors.append(f'slide {i + 1} {ct}: no montage in {folder}')
                    continue
                if re.findall(r'\d+', old)[:1] != re.findall(r'\d+', new.stem)[:1]:
                    errors.append(f'slide {i + 1} {ct}: first FOV differs ({old} vs {new.name})')
                plan.append((i, pic, new))
        else:                                             # plot slide: same-named panel
            for pic in pics:
                new = plot_source(descr_of(pic).get('descr') or '')
                if not new.is_file():
                    errors.append(f'slide {i + 1}: missing {new}')
                    continue
                plan.append((i, pic, new))
    return plan, errors


def swap(slide, pic, new):
    """Point the picture at the new image; aspect-fit in the old box, centred; drop the old rel."""
    part = slide.part
    blip = pic._element.xpath('./p:blipFill/a:blip')[0]
    old_rid = blip.get(qn('r:embed'))
    _, rid = part.get_or_add_image_part(str(new))
    blip.set(qn('r:embed'), rid)
    # python-pptx's drop_rel only counts r:id refs (not a:blip/@r:embed) -> count embeds ourselves.
    if rid != old_rid and old_rid not in slide._element.xpath('.//@r:embed | .//@r:id | .//@r:link'):
        part._rels.pop(old_rid)
    descr_of(pic).set('descr', new.name)
    with bpd.Image.open(new) as im:
        ar = im.size[0] / im.size[1]
    L, T, W, H = pic.left, pic.top, pic.width, pic.height
    w, h = (W, W / ar) if ar >= W / H else (H * ar, H)
    pic.left, pic.top = Emu(int(L + (W - w) / 2)), Emu(int(T + (H - h) / 2))
    pic.width, pic.height = Emu(int(w)), Emu(int(h))


class Shim:
    """tight_grid() only needs .slide(title); add blank slides to THIS deck."""

    def __init__(self, prs, layout):
        self.prs, self.layout, self.new = prs, layout, []

    def slide(self, title):
        s = self.prs.slides.add_slide(self.layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        bpd.add_text(s, 0.3, 0.1, bpd.SLIDE_W - 0.6, 0.55, title, 26, bold=True, align=bad.PP_ALIGN.CENTER)
        self.new.append(s)
        return s


def add_profile_slides(prs):
    layout = next(sl.slide_layout for sl in prs.slides if sl.slide_layout.name == 'Blank')
    deck = Shim(prs, layout)
    P = lambda n: PANELS / 'profiles' / n  # noqa: E731
    V = lambda n: PANELS / 'violins' / n  # noqa: E731
    tcols = [lab for _, lab, _, _ in bad.DATASETS]
    for tag, lab, _, _ in bad.DATASETS:
        bad.tight_grid(deck, f'Actin Profile and Gradient {bad.DASH} {lab}',
                       [P(f'both_catfmc_profile_{tag}_{t}.png') for t in bad.TPS] +
                       [P(f'both_catfmc_grad_{tag}_{t}.png') for t in bad.TPS], bad.TP_LBL)
    bad.tight_grid(deck, 'Actin Profile Time Course',
                   [P(f'tc_{ct}_{kind}_{tag}.png') for tag, _, _, _ in bad.DATASETS
                    for kind in ('profile', 'grad') for ct in ('CAT', 'FMC')],
                   ['CAT Profile', 'FMC Profile', 'CAT Gradient', 'FMC Gradient'], tcols)
    bad.tight_grid(deck, 'Actin by Zone: Normalized by MFI over Full Synapse',
                   [V(f'zones_actin_{tag}_{t}.png') for tag, _, _, _ in bad.DATASETS for t in bad.TPS],
                   bad.TP_LBL, tcols, key=True)
    bad.tight_grid(deck, 'Actin Zone MFI (Raw)',
                   [V(f'zones_raw_actin_{tag}_{t}.png') for tag, _, _, _ in bad.DATASETS for t in bad.TPS],
                   bad.TP_LBL, tcols, key=True)
    for s in deck.new:                                    # text above pictures
        tree = s.shapes._spTree
        for sh in list(s.shapes):
            if sh.has_text_frame:
                tree.remove(sh._element); tree.append(sh._element)
    lst = prs.slides._sldIdLst
    ids = list(lst)[-len(deck.new):]
    for k, el in enumerate(ids):                          # move right after slide INSERT_AFTER
        lst.remove(el); lst.insert(INSERT_AFTER + k, el)
    return [title_of(s) for s in deck.new]


def nonpic_xml(slide):
    out = []
    for el in slide.shapes._spTree.iterchildren():
        if el.tag == qn('p:pic'):
            continue
        out.append(etree.tostring(copy.deepcopy(el)))
    return out


def verify(expected, n_new, src_sha, shift=None, index_map=None):
    """expected: list of (new_slide_idx0, picture shape id, source path, old blob sha).
    n_new = slides added in total; shift = how many of them sit before original slide INSERT_AFTER+1
    (default all of them; the rest are appended at the end). index_map: {old idx0: new idx0} for
    decks whose original slides were also reordered (overrides INSERT_AFTER/shift)."""
    shift = n_new if shift is None else shift
    src, out = Presentation(str(SRC)), Presentation(str(OUT))
    fails = []
    n_old = len(src.slides)
    if len(out.slides) != n_old + n_new:
        fails.append(f'slide count {len(out.slides)} != {n_old + n_new}')
    for i, s in enumerate(src.slides):
        if index_map is not None and i not in index_map:
            continue                                      # slide deliberately removed
        j = index_map[i] if index_map is not None else (i if i < INSERT_AFTER else i + shift)
        t = out.slides[j]
        if nonpic_xml(s) != nonpic_xml(t):
            fails.append(f'slide {i + 1} -> {j + 1}: non-picture XML changed')
        if (s.has_notes_slide and s.notes_slide.notes_text_frame.text) != \
           (t.has_notes_slide and t.notes_slide.notes_text_frame.text):
            fails.append(f'slide {i + 1}: notes changed')
        op = {sh.shape_id: sh for sh in s.shapes if sh.shape_type == 13}
        npics = {sh.shape_id: sh for sh in t.shapes if sh.shape_type == 13}
        if op.keys() != npics.keys():
            fails.append(f'slide {i + 1}: picture set changed')
        for sid, o in op.items():
            n = npics.get(sid)
            if n is None:
                continue
            tol = 2
            if not (n.left >= o.left - tol and n.top >= o.top - tol and n.left + n.width <= o.left + o.width + tol
                    and n.top + n.height <= o.top + o.height + tol):
                fails.append(f'slide {i + 1} pic {sid}: outside old bbox')
    by_key = {(j, sid): (p, old) for j, sid, p, old in expected}
    for (j, sid), (p, old) in by_key.items():
        pic = next(sh for sh in out.slides[j].shapes if sh.shape_id == sid)
        b = sha(pic.image.blob)
        if b != sha(Path(p).read_bytes()):
            fails.append(f'slide {j + 1} pic {sid}: blob != {p}')
        if b == old:
            fails.append(f'slide {j + 1} pic {sid}: unchanged blob')
    for s in out.slides:                                  # every image rel is referenced
        refs = set(s._element.xpath('.//a:blip/@r:embed'))
        for rid, rel in s.part.rels.items():
            if rel.reltype.endswith('/image') and rid not in refs:
                fails.append(f'orphan image rel {rid} on slide {title_of(s)!r}')
    if sha(SRC.read_bytes()) != src_sha:
        fails.append('SOURCE CHANGED')
    return fails


def main():
    if OUT.exists() or (OUT.parent / ('~$' + OUT.name)).exists():
        print(f'[abort] {OUT} (or its lock) already exists'); return 1
    src_sha = sha(SRC.read_bytes())
    if TWIN.is_file() and sha(TWIN.read_bytes()) != src_sha:
        print(f'[abort] {SRC} and {TWIN} differ'); return 1
    prs = Presentation(str(SRC))
    plan, errors = resolve(prs)
    if errors:
        print('[abort] unresolved:\n  ' + '\n  '.join(errors)); return 1
    print(f'{len(plan)} pictures resolved')
    expected = []
    for i, pic, new in plan:
        old_sha = sha(pic.image.blob)
        print(f'  slide {i + 1:2d} {descr_of(pic).get("descr"):70s} -> {new}')
        swap(prs.slides[i], pic, new)
        j = i if i < INSERT_AFTER else None               # index after insertion, fixed below
        expected.append((i, pic.shape_id, str(new), old_sha, j))
    titles = add_profile_slides(prs)
    n_new = len(titles)
    prs.save(str(OUT))
    print(f'saved {OUT} ({len(prs.slides)} slides); new slides {INSERT_AFTER + 1}-{INSERT_AFTER + n_new}: {titles}')
    exp = [(i if i < INSERT_AFTER else i + n_new, sid, p, old) for i, sid, p, old, _ in expected]
    fails = verify(exp, n_new, src_sha)
    print('VERIFY: ' + ('PASS' if not fails else 'FAIL\n  ' + '\n  '.join(fails)))
    return 0 if not fails else 2


if __name__ == '__main__':
    sys.exit(main())
