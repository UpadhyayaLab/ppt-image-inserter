#!/usr/bin/env python3
r"""montage_scale.py  Check (and optionally fix) that CAT-vs-FMC montage slides show their
montages at ONE physical scale (2026-10-17: decks that aspect-fit each montage to its box
shrank CAT and FMC -- and each timepoint -- by different factors).

A montage slide = a slide with >= 2 pictures and text boxes 'CAT' and 'FMC...'. Montages of
one view share a title stem (the title before ' \u2014 '). For every picture the displayed
scale is source pixels / displayed inches (crop-aware). Within one view all pictures should
share one px/in -- valid because each view's montage PNGs are rendered at one px/um
(set_figure_for_physical_scale; check with PPT_image_inserter check_scalebar_pixel_widths.py).

  python montage_scale.py --check <deck.pptx> [<deck.pptx> ...]   # read-only report
  python montage_scale.py --check-dir K:\FF\PPT\PPT_autogeneration\CART
  python montage_scale.py --fix <deck.pptx>    # writes <deck>_onescale.pptx next to it

--fix (for decks that cannot be rebuilt, e.g. FF-edited ones): in each view, every picture is
shrunk (never enlarged) to the view's highest px/in, centered where it was. The source deck
is never modified.
"""
from __future__ import annotations

import argparse
import io
import sys
from collections import defaultdict
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_IN = 914400
DASH = '\u2014'
TOL = 0.02          # report a view when its px/in spread exceeds 2 %


def slide_title(slide):
    texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    for t in texts:
        if DASH in t:
            return t
    return texts[0] if texts else ''


def is_montage_slide(slide):
    texts = {sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame}
    pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    return len(pics) >= 2 and 'CAT' in texts and any(t.startswith('FMC') for t in texts)


def picture_ppi(pic):
    """Source px per displayed inch along x, accounting for cropping."""
    with Image.open(io.BytesIO(pic.image.blob)) as im:
        w_px = im.size[0]
    shown_px = w_px * (1 - pic.crop_left - pic.crop_right)
    return shown_px / (pic.width / EMU_IN)


def montage_pictures(prs):
    """-> {view stem: [(slide index, picture, px/in), ...]}"""
    views = defaultdict(list)
    for i, s in enumerate(prs.slides, 1):
        if not is_montage_slide(s):
            continue
        stem = slide_title(s).split(f' {DASH} ')[0]
        for sh in s.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                views[stem].append((i, sh, picture_ppi(sh)))
    return views


def check(deck):
    views = montage_pictures(Presentation(str(deck)))
    rows = []
    for stem, items in views.items():
        ppis = [p for _, _, p in items]
        lo, hi = min(ppis), max(ppis)
        n_slides = len({i for i, _, _ in items})
        # within-slide: CAT and FMC on the same slide at different scales
        by_slide = defaultdict(list)
        for i, _, p in items:
            by_slide[i].append(p)
        worst_slide = max((max(v) / min(v) - 1) for v in by_slide.values())
        rows.append((stem, n_slides, lo, hi, hi / lo - 1, worst_slide))
    return rows


def report(deck, rows):
    bad_rows = [r for r in rows if r[4] > TOL]
    print(f'\n{deck}  ({len(rows)} montage view(s), {len(bad_rows)} not at one scale)')
    for stem, n, lo, hi, spread, worst in rows:
        flag = 'OFF' if spread > TOL else 'ok '
        print(f'  {flag} {stem[:58]:58s} slides {n:3d}  px/in {lo:7.1f}-{hi:7.1f}  '
              f'spread {100 * spread:5.1f}%  worst CAT/FMC on one slide {100 * worst:5.1f}%')
    return len(bad_rows)


def fix(deck):
    prs = Presentation(str(deck))
    n_changed = 0
    for items in montage_pictures(prs).values():
        target = max(p for _, _, p in items)
        for _, pic, p in items:
            if p >= target * (1 - 1e-3):
                continue
            f = p / target                       # < 1: shrink about the center
            cx, cy = pic.left + pic.width / 2, pic.top + pic.height / 2
            w, h = round(pic.width * f), round(pic.height * f)
            pic.left, pic.top, pic.width, pic.height = round(cx - w / 2), round(cy - h / 2), w, h
            n_changed += 1
    out = deck.with_name(deck.stem + '_onescale.pptx')
    prs.save(str(out))
    print(f'saved {out} ({n_changed} picture(s) rescaled)')
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', nargs='*', default=[], help='decks to report (read-only)')
    ap.add_argument('--check-dir', help='report every .pptx under this folder (read-only)')
    ap.add_argument('--fix', help='deck to copy to <name>_onescale.pptx at one scale per view')
    args = ap.parse_args()
    decks = [Path(d) for d in args.check]
    if args.check_dir:
        decks += sorted(p for p in Path(args.check_dir).rglob('*.pptx') if not p.name.startswith('~$'))
    n_off = 0
    for d in decks:
        try:
            n_off += report(d, check(d))
        except Exception as e:                   # noqa: BLE001  (one unreadable deck must not stop the audit)
            print(f'\n{d}  [error] {e}')
    if args.fix:
        out = fix(Path(args.fix))
        n_off += report(out, check(out))
    return 1 if n_off else 0


if __name__ == '__main__':
    sys.exit(main())
