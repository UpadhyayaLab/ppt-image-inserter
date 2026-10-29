"""Cathepsin B | Granzyme B side-by-side deck: the same polarization metrics for the three
CTSB/MT experiments (Nov 27 2023, Jun 20 2024, Jun 24 2024) and Kiet's Oct 8, 2026 GzmB set,
in date order, one panel per experiment, never pooled (FF + Kiet 2026-10-29).

Sources (no figure is drawn here):
  violins    <root>\\gzmb_vs_ctsb_violins_{qc,pos}_<newest>\\violins  tests/run_gzmb_vs_ctsb_violins.m
             (the GzmB deck's violin style, one y-range and one width per timepoint for all panels)
  fractions  GzmB <Oct 8 root>\\compiled_<newest>\\polarization        d20261008_gzmb_compilation_config
             CTSB <CTSB root>\\compiled_ctsb_vs_gzmb_<newest>\\polarization
             d20231127_20240620_20240624_ctsb_vs_gzmb_compilation_config (July 2026 CTSB results
             compiled with the GzmB cell QC, cutoffs and tests)

  python build_gzmb_vs_ctsb_deck.py --export <dir>
-> K:\\FF\\PPT\\PPT_autogeneration\\CART\\GzmB\\CART_GzmB_vs_CTSB_<yyyymmdd>.pptx
"""
from __future__ import annotations

import argparse
import datetime
import sys
from pathlib import Path

from PIL import Image
from pptx.enum.text import PP_ALIGN

import build_gzmb_decks as bgd          # also puts CART_pMLC_ppt / CART_actin_ppt on sys.path
import build_pmlc_profiles_deck as bpd

MU, DASH, SUB50, SUB75 = bgd.MU, bgd.DASH, bgd.SUB50, bgd.SUB75
GZMB_ROOT = bgd.ROOT
CTSB_ROOT = Path(r'J:\FF\fixed_cell\CAR_TCell\results_compiled\MT_CATB_20231127_20240620_20240624')
OUT = bgd.OUT_DIR / f'CART_GzmB_vs_CTSB_{datetime.date.today():%Y%m%d}.pptx'

# Date order (FF 2026-10-29): (protein, experiment name, violin tag, which root)
EXPS = [('Cathepsin B', 'Nov 27, 2023', 'Nov_27', 'ctsb'),
        ('Cathepsin B', 'Jun 20, 2024', 'Jun_20', 'ctsb'),
        ('Cathepsin B', 'Jun 24, 2024', 'Jun_24', 'ctsb'),
        ('Granzyme B', 'Oct 8, 2026', 'Oct_8', 'gzmb')]

POS = 'GzmB+ = whole-cell GzmB total signal > 10\u2076; CTSB+ = CTSB peak signal > 20.'
CENT = 'Centrosome = brightest point of the smoothed MT stack within the actin MIP outline.'
VIOLIN_NOTE = ('White dot: median; bar: interquartile range. Stars: Wilcoxon rank-sum test, '
               'CAT vs FMC, Benjamini-Hochberg corrected.')

# (title, violin group, GzmB field, CTSB field, cutoffs, note)
METRICS = [
    ('Centrosome\u2013Synapse Distance', 'qc', 'centrosome_center_z_cell_bottom_distance',
     'centrosome_center_z_cell_bottom_distance', [2, 3], CENT),
    (f'Granule {SUB50} ({MU}m from Cell Bottom) {DASH} Positive Cells Only', 'pos',
     'GzmB_z50_cell_bottom_distance', 'CathepsinB_z50_rel_cell_bottom', [2, 3], POS),
    (f'Granule {SUB75} ({MU}m from Cell Bottom) {DASH} Positive Cells Only', 'pos',
     'GzmB_z75_cell_bottom_distance', 'CathepsinB_z75_rel_cell_bottom', [3], POS),
]

MISSING: list[Path] = []


def check(p: Path) -> Path:
    if not p.is_file():
        MISSING.append(p)
    return p


def aspect(p: Path) -> float:
    if not p.is_file():
        return 1.0
    with Image.open(p) as im:
        return im.size[0] / im.size[1]


def note_text(s, text):
    tb = bpd.add_text(s, 0.15, bpd.SLIDE_H - 0.5, bpd.SLIDE_W - 0.3, 0.45, text, 14,
                      align=PP_ALIGN.CENTER)
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.italic = True


def header(s, left, width, text, top, size=18):
    bpd.add_text(s, left, top, width, 0.4, text, size, bold=True, align=PP_ALIGN.CENTER)


def row_slide(deck, title, pngs, note):
    """One row in date order at one height, so every timepoint gets the same width (the
    violins are rendered at one width per timepoint). Protein header over its experiments,
    the experiment date over each panel."""
    s = deck.slide(title)
    gap, group_gap = 0.08, 0.3
    top, bottom = 1.55, bpd.SLIDE_H - 0.6
    asp = [aspect(p) for p in pngs]
    n_groups = len({e[0] for e in EXPS})
    free_w = bpd.SLIDE_W - 0.3 - gap * (len(pngs) - n_groups) - group_gap * (n_groups - 1)
    h = min(bottom - top, free_w / sum(asp))
    widths = [a * h for a in asp]
    x = (bpd.SLIDE_W - (sum(widths) + gap * (len(pngs) - n_groups) + group_gap * (n_groups - 1))) / 2
    y = top + (bottom - top - h) / 2
    groups: dict[str, list[float]] = {}
    for (protein, name, _, _), png, w in zip(EXPS, pngs, widths):
        if protein in groups:
            x += gap
        elif groups:
            x += group_gap
        groups.setdefault(protein, [x, x])[1] = x + w
        header(s, x, w, name, y - 0.4, 16)
        bpd.fit_in_box(s, png, x, y, w, h)
        x += w
    for protein, (x0, x1) in groups.items():
        header(s, x0, x1 - x0, protein, y - 0.8, 20)
    note_text(s, note)


def quad_slide(deck, title, cells, note):
    """2 x 2 in date order (left to right, top to bottom)."""
    s = deck.slide(title)
    top, bottom, lab = 0.75, bpd.SLIDE_H - 0.55, 0.36
    cw = (bpd.SLIDE_W - 0.3) / 2
    ch = (bottom - top) / 2
    for i, (label, png) in enumerate(cells):
        r, c = divmod(i, 2)
        x, y = 0.15 + c * cw, top + r * ch
        header(s, x, cw, label, y)
        bpd.fit_in_box(s, png, x + 0.05, y + lab, cw - 0.1, ch - lab - 0.04)
    note_text(s, note)


def newest(root: Path, pattern: str) -> Path:
    d = bpd.newest_run(root, pattern)
    if d is None:
        sys.exit(f'[error] no {pattern} under {root}')
    return d


def build(export):
    comp = {'gzmb': newest(GZMB_ROOT, 'compiled_2*'),
            'ctsb': newest(CTSB_ROOT, 'compiled_ctsb_vs_gzmb_2*')}
    vio = {(r, g): newest(root, f'gzmb_vs_ctsb_violins_{g}_2*') / 'violins'
           for r, root in (('gzmb', GZMB_ROOT), ('ctsb', CTSB_ROOT)) for g in ('qc', 'pos')}
    print('\n'.join(f'{k}: {v}' for k, v in {**comp, **vio}.items()))
    bpd.OUT = OUT
    deck = bgd.OvDeck()

    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.titles.append('Title')
    bpd.add_text(s, 0.83, 2.2, 11.67, 1.2, 'Cathepsin B and Granzyme B: Polarization to the Synapse',
                 40, bold=True, align=PP_ALIGN.CENTER)
    bpd.add_text(s, 0.83, 3.55, 11.67, 0.5, f'CAR T Cells, CAT vs FMC {DASH} One Panel per Experiment',
                 24, align=PP_ALIGN.CENTER)
    today = datetime.date.today()
    bpd.add_text(s, 4.37, 4.4, 4.6, 0.5, f'{today:%B} {today.day}, {today.year}', 20,
                 align=PP_ALIGN.CENTER)

    s = deck.slide('Data and Cell Selection')
    bpd.add_text(s, 0.8, 1.0, bpd.SLIDE_W - 1.6, 6.2, ['\u2022 ' + b for b in (
        ('Cathepsin B: Nov 27, 2023 (15 min), Jun 20, 2024 and Jun 24, 2024 (5, 15 min); '
         'processed Jul 2, 2026.'),
        'Granzyme B: Oct 8, 2026, CAT vs FMC at 5, 10, 15 min; processed Oct 22, 2026.',
        (f'Same cell selection for every experiment: synapse slice \u2264 10 and cell height '
         f'\u2265 2 {MU}m. Experiments are shown separately, never pooled.'),
        f'Centrosome: {CENT[len("Centrosome = "):-1]}; height above the synapse slice. Same code for both proteins.',
        (f'{SUB50} / {SUB75}: height above the synapse slice below which 50% / 75% of the cell\u2019s '
         f'granule signal lies; same computation for both proteins. Positive cells only: {POS}'),
        ('Not compared yet (the July Cathepsin B values use older definitions; needs reprocessing): '
         f'Z-COF and the fraction of signal within 1/2/3 {MU}m of the centrosome.'),
    )], 20)

    for title, grp, gf, cf, cutoffs, note in METRICS:
        fld = {'gzmb': gf, 'ctsb': cf}
        row_slide(deck, title,
                  [check(vio[(r, grp)] / f'puncta_{fld[r]}_{tag}.png') for _, _, tag, r in EXPS],
                  f'{note} {VIOLIN_NOTE}')
        stem = title.split(f' {DASH} ')[0].split(' (')[0]
        gate = f' {DASH} Positive Cells Only' if grp == 'pos' else ''
        for t in cutoffs:
            cells = [(f'{protein} {DASH} {name}',
                      check(comp[r] / 'polarization' / name.replace(' ', '_') / f'{fld[r]}_pol_thresh{t}.png'))
                     for protein, name, _, r in EXPS]
            quad_slide(deck, f'{stem}: Fraction Polarized (\u2264 {t} {MU}m){gate}', cells,
                       f'{note} {bgd.FRACTION_NOTE}')

    deck.save()
    print(f'saved {OUT} ({len(deck.prs.slides)} slides)')
    if export:
        bpd.export_pngs(Path(export), deck.titles)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--export', help='also export one PNG per slide (named by title) under this dir')
    args = ap.parse_args()
    build(args.export)
    if MISSING:
        print(f'\n[MISSING] {len(MISSING)} panel(s):')
        for m in MISSING:
            print(f'  {m}')
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
