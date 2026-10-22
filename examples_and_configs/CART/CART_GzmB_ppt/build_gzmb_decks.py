#!/usr/bin/env python3
r"""build_gzmb_decks.py  Review decks for Kiet's 20261008 fixed CAR-T granzyme B set
(CAT / FMC x 5 / 10 / 15 min; actin C1, MT C2, GzmB C3; IRM C4 left out until its
identity is confirmed). Same slide style and metric sets as the recent CART decks:
  - Centrosome / GzmB block = CART_summary_MT_CatB.pptx order, GzmB in place of CatB
  - actin block       = CART_actin_ppt/build_actin_decks.py SUMMARY_OLD + SUMMARY_NEW
  - montages          = build_actin_decks.build_montages (first page, CAT left / FMC right)

Decks (K:\FF\PPT\PPT_autogeneration\CART\GzmB\CART_GzmB_<deck>_20261008_<yyyymmdd>.pptx):
  summary    cell counts, centrosome / GzmB, actin profiles + zones, actin metrics
  zones      GzmB / MT central-peripheral-distal zones, centrosome + GzmB-COF position (c/p/dSMAC)
  threshold  GzmB total / peak signal (ungated) + GzmB+ fraction, to set the GzmB+ cutoff
  montages   actin, MT and GzmB per-cell QC montages
  overview   CART_GzmB_MT_20261008_<yyyymmdd>.pptx: format + content of FF's hand-made
             CART_CatB_MT_20260702.pptx (montages, then one compile grid per slide) plus GzmB
             amount (log), GzmB+ fraction and zones (FF 2026-10-20)

Inputs (nothing is plotted here):
  <root>\compiled_2*\  (newest)       <- run/run_gzmb_20261008.m (compile_cart d20261008_gzmb)
  <root>\actin_extended_panels_*\     <- tests/run_gzmb_20261008_extended_panels.m
  <root>\<COND>\converted\cropped\split_channels\prog_fixed_cells\...\montages_filtered\

Run:  python build_gzmb_decks.py [--deck summary|zones|threshold|montages|all] [--export <dir>]
Exits 1 if any panel is missing (each missing file is listed).
"""
from __future__ import annotations

import argparse
import datetime
import glob
import math
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / 'CART_pMLC_ppt'))
sys.path.insert(0, str(HERE.parent / 'CART_actin_ppt'))
import build_pmlc_profiles_deck as bpd
import build_actin_decks as bad          # also: .mont = the established montage-slide helpers
from pptx.enum.text import PP_ALIGN

ROOT = Path(r'\\10.225.56.176\ULab_data\User_data\Kiet\Polarization CAR T'
            r'\20261008_,fix_CART_granzyme_5_10_15')
OUT_DIR = Path(r'K:\FF\PPT\PPT_autogeneration\CART\GzmB')
TODAY = f'{datetime.date.today():%Y%m%d}'
DASH, MINUS, MU = '\u2014', '\u2212', '\u03bc'
SUB50, SUB75, SUB1 = 'Z\u2085\u2080', 'Z\u2087\u2085', '\u2081'

TPS = ['5min', '10min', '15min']
TP_LBL = ['5 min', '10 min', '15 min']
TAG = 'Oct_8'                      # experiment tag used in the extended-panel filenames
TS = 'timecourse_scatter_plots'
CVF = 'CAT_vs_FMC_by_timepoint'

# --- Centrosome / GzmB: CART_summary_MT_CatB.pptx order, CatB -> GzmB -------------------------
GZMB_BLOCK = [
    ('Centrosome-Synapse Distance', f'{TS}/centrosome_center_z_cell_bottom_distance_grid.png'),
    ('Centrosome-Synapse Distance: Fraction Polarized', 'centrosome_center_z_cell_bottom_distance_polarization_grid.png'),
    ('GzmB+ Fraction at Three Total-Signal Cutoffs', 'GzmB_total_sig_polarization_grid.png'),
    ('GzmB Synapse Inner/Outer Ratio (50% Eff Rad Thresh)', f'{TS}/GzmB_synapse_inner_outer_ratio_grid.png'),
    ('GzmB Synapse (3-Slice MIP) Inner/Outer Ratio (50% Eff Rad Thresh)',
     f'{TS}/GzmB_synapse_inner_outer_ratio_3mip_grid.png'),
    ('GzmB Synapse MFI', f'{TS}/GzmB_synapse_MFI_grid.png'),
    # 3-slice MIP total, the plane of the synapse MFI and profiles (FF 2026-10-21)
    ('GzmB Synapse Total Signal', f'{TS}/GzmB_synapse_total_sig_3mip_grid.png'),
    (f'GzmB Synapse Total Signal {DASH} CAT vs FMC', f'{CVF}/GzmB_synapse_total_sig_3mip_log_grid.png'),
    ('GzmB Total Signal', f'{TS}/GzmB_total_sig_grid.png'),
    (f'GzmB Total Signal {DASH} CAT vs FMC', f'{CVF}/GzmB_total_sig_log_grid.png'),
    (f'GzmB Total Signal CDF {DASH} CAT vs FMC', f'{CVF}/GzmB_total_sig_cdf_grid.png'),
    (f'GzmB Fraction Within 1 {MU}m of Centrosome', f'{TS}/GzmB_frac_around_cent_1um_grid.png'),
    (f'GzmB Fraction Within 2 {MU}m of Centrosome', f'{TS}/GzmB_frac_around_cent_2um_grid.png'),
    (f'GzmB Fraction Within 3 {MU}m of Centrosome', f'{TS}/GzmB_frac_around_cent_3um_grid.png'),
    (f'GzmB {SUB50} ({MU}m from Cell Bottom)', f'{TS}/GzmB_z50_cell_bottom_distance_grid.png'),
    (f'GzmB {SUB50} (Normalized by Cell Height)', f'{TS}/GzmB_z50_cell_bottom_distance_norm_cell_height_grid.png'),
    (f'GzmB {SUB50}: Fraction Polarized', 'GzmB_z50_cell_bottom_distance_polarization_grid.png'),
    (f'GzmB {SUB75} ({MU}m from Cell Bottom)', f'{TS}/GzmB_z75_cell_bottom_distance_grid.png'),
    (f'GzmB {SUB75} (Normalized by Cell Height)', f'{TS}/GzmB_z75_cell_bottom_distance_norm_cell_height_grid.png'),
    (f'GzmB {SUB75}: Fraction Polarized', 'GzmB_z75_cell_bottom_distance_polarization_grid.png'),
]


def actin_block():
    """build_actin_decks SUMMARY_OLD + SUMMARY_NEW, minus IRM (C4 unconfirmed). This compile
    stores the actin radial profile as actin_bottom_slice_rad_profile (no _auc1, no 3-slice)."""
    out = []
    for title, rel in bad.SUMMARY_OLD + bad.SUMMARY_NEW:
        if 'IRM' in rel:
            continue
        if rel.endswith('actin_bottom_slice_rad_profile_auc1_all_cells_with_average_grid.png'):
            title = 'Actin Synapse Radial Profile (All Cells + Avg)'
            rel = f'{CVF}/actin_bottom_slice_rad_profile_all_cells_with_average_grid.png'
        elif 'rad_profile_auc1' in rel:
            continue                    # 3-slice MIP radial profile is not in this compile
        out.append((title, rel))
    return out


ZONES = ['Central', 'Peripheral', 'Distal']
ZONE_BLOCK = (
    [(f'GzmB Signal Fraction in {z} Zone', f'{TS}/GzmB_synapse_{z.lower()}_sigfrac_grid.png') for z in ZONES]
    + [(f'GzmB {z} MFI / Synapse Mean', f'{TS}/GzmB_synapse_{z.lower()}_MFI_norm_grid.png') for z in ZONES]
    + [(f'MT Signal Fraction in {z} Zone', f'{TS}/MT_synapse_{z.lower()}_sigfrac_grid.png') for z in ZONES]
    + [(f'MT {z} MFI / Synapse Mean', f'{TS}/MT_synapse_{z.lower()}_MFI_norm_grid.png') for z in ZONES]
    + [('Centrosome XY Offset / Synapse Radius', f'{TS}/centrosome_xy_offset_norm_grid.png'),
       ('Centrosome XY Offset / Synapse Radius: Fraction Centered', 'centrosome_xy_offset_norm_polarization_grid.png'),
       (f'Centrosome XY Offset from Synapse Center ({MU}m)', f'{TS}/centrosome_xy_offset_um_grid.png'),
       ('Centrosome Synapse Zone (1 Central, 2 Peripheral, 3 Distal, 0 Outside)', f'{TS}/centrosome_synapse_zone_grid.png'),
       ('GzmB COF XY Offset / Synapse Radius', f'{TS}/GzmB_COF_xy_offset_norm_grid.png'),
       ('GzmB COF Synapse Zone (1 Central, 2 Peripheral, 3 Distal, 0 Outside)', f'{TS}/GzmB_COF_synapse_zone_grid.png')]
)

THRESHOLD_BLOCK = [
    (f'GzmB Total Signal CDF {DASH} CAT vs FMC (All Cells)', f'{CVF}/GzmB_total_sig_cdf_grid.png'),
    (f'GzmB Total Signal CDF {DASH} Time Course (All Cells)', 'time_course_by_cell_type/GzmB_total_sig_cdf_grid.png'),
    (f'GzmB Total Signal {DASH} CAT vs FMC (All Cells, Ungated)', f'{CVF}/GzmB_total_sig_log_grid.png'),
    (f'GzmB Total Signal {DASH} Time Course (All Cells, Ungated)',
     'time_course_by_cell_type/GzmB_total_sig_log_grid.png'),
    ('GzmB Total Signal (All Cells, Ungated)', f'{TS}/GzmB_total_sig_grid.png'),
    ('GzmB+ Fraction at Three Total-Signal Cutoffs', 'GzmB_total_sig_polarization_grid.png'),
    (f'GzmB Peak Signal {DASH} CAT vs FMC (All Cells, Ungated)', f'{CVF}/GzmB_peak_sig_log_grid.png'),
    ('GzmB Peak Signal (All Cells, Ungated)', f'{TS}/GzmB_peak_sig_grid.png'),
]

# (title stem, prog_fixed_cells sub-folder, scale group). Every slide of one scale group is
# placed at ONE pixels-per-inch (insert_CART_CatB_MT_actin_nuc_montages_slides.py rule:
# pinned to the largest montage in the group), so CAT, FMC and all timepoints share one
# physical scale. The per-cell renders behind every group use set_figure_for_physical_scale
# (one project PPUM); 'centrosome' since 2026-10-17 (plot_centrosome_xzmip.m).
MONTAGES = [
    # Every synapse view names its plane: 'Synapse Slice' (1 plane) vs '3-Slice MIP' (FF 2026-10-20)
    ('Actin, Synapse Slice', r'actin\synapse\inner_outer\1slice_combined', 'actin_syn_combined'),
    ('Actin XZ MIP', r'actin\xz_mip', 'xz'),
    ('MT, Synapse Slice', r'MT\synapse\1slice\raw', 'synapse'),
    ('MT XZ MIP with Centrosome', r'MT\centrosome', 'centrosome'),
    ('MT XZ MIP', r'MT\xz_mip', 'xz'),
    ('GzmB, Synapse Slice', r'GzmB\synapse\1slice\raw', 'synapse'),
    ('GzmB, Synapse Slice (Composite)', r'GzmB\synapse\composite\1slice', 'synapse'),
    ('GzmB XZ MIP', r'GzmB\xz_mip', 'xz'),
    # a list = contrast variants shown on consecutive slides for each timepoint
    [('GzmB XZ MIP (Fixed Contrast 10–100)', r'GzmB\xz_mip\fixed_con_10_100', 'xz'),
     ('GzmB XZ MIP (Fixed Contrast 10–500)', r'GzmB\xz_mip\fixed_con_10_500', 'xz')],
]

MISSING: list[Path] = []


def check(p):
    p = Path(p) if p else None
    if p is None or not p.is_file():
        MISSING.append(p)
    return p


def newest(pattern):
    hits = sorted(p for p in glob.glob(str(ROOT / pattern)) if Path(p).is_dir())
    return Path(hits[-1]) if hits else None


def start(deck_name):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    bpd.OUT = OUT_DIR / f'CART_GzmB_{deck_name}_20261008_{TODAY}.pptx'
    return bpd.Deck()


def finish(deck, export, sub):
    deck.save()
    print(f'saved {bpd.OUT} ({len(deck.prs.slides)} slides)')
    if export:
        bpd.export_pngs(Path(export) / sub, deck.titles)


# GzmB fields the compile does NOT gate (gzmb_raw_signal in d20261008_gzmb_compilation_config):
# every cell. Every other GzmB_* field is computed on GzmB+ cells only (GzmB_total_sig > 1e6).
GZMB_UNGATED = {
    'GzmB_peak_sig', 'GzmB_total_sig', 'GzmB_synapse_MFI', 'GzmB_synapse_MFI_3mip',
    'GzmB_synapse_total_sig', 'GzmB_synapse_total_sig_3mip',
    'GzmB_synapse_inner_mask_MFI', 'GzmB_synapse_inner_mask_MFI_3mip',
    'GzmB_synapse_outer_mask_MFI', 'GzmB_synapse_outer_mask_MFI_3mip'}


def gated_title(title, rel):
    """Say on the slide when a panel shows GzmB+ cells only (review 2026-10-17: the deck mixed
    all-cell and GzmB+-only panels with nothing on the slide telling them apart)."""
    field = Path(rel).name
    for suffix in ('_polarization_grid.png', '_cdf_grid.png', '_log_grid.png', '_grid.png'):
        if field.endswith(suffix):
            field = field[:-len(suffix)]
            break
    if field.startswith('GzmB_') and field not in GZMB_UNGATED:
        return f'{title} {DASH} GzmB+ Cells Only'
    return title


def singles(deck, compile_dir, entries):
    for title, rel in entries:
        bad.single(deck, gated_title(title, rel), check(compile_dir / 'grid_panels' / rel))


def build_summary(compile_dir, panels, export):
    deck = start('summary')
    bad.single(deck, 'Cells per Condition', check(compile_dir / 'cell_counts_barplot.png'))
    singles(deck, compile_dir, GZMB_BLOCK)

    P = lambda n: check(panels / 'profiles' / n)
    V = lambda n: check(panels / 'violins' / n)
    bad.tight_grid(deck, f'Actin Profile and Gradient {DASH} Oct 8, 2026',
                   [P(f'both_catfmc_profile_{TAG}_{t}.png') for t in TPS] +
                   [P(f'both_catfmc_grad_{TAG}_{t}.png') for t in TPS], TP_LBL)
    bad.tight_grid(deck, 'Actin Profile Time Course',
                   [P(f'tc_{ct}_{kind}_{TAG}.png') for kind in ('profile', 'grad') for ct in ('CAT', 'FMC')],
                   ['CAT', 'FMC'])   # 2 x 2 (profile row, gradient row): four in a row were tiny
    bad.tight_grid(deck, f'FMC {MINUS} CAT Actin Profile',
                   [P(f'diff_both_profile_{t}.png') for t in TPS] +
                   [P(f'diff_both_grad_{t}.png') for t in TPS], TP_LBL)
    bad.tight_grid(deck, 'Actin by Zone: Normalized by MFI over Full Synapse',
                   [V(f'zones_actin_{TAG}_{t}.png') for t in TPS], TP_LBL, key=True)
    bad.tight_grid(deck, 'Actin Zone MFI (Raw)',
                   [V(f'zones_raw_actin_{TAG}_{t}.png') for t in TPS], TP_LBL, key=True)
    bad.single(deck, 'Actin Synapse MFI', V(f'synapse_mfi_{TAG}.png'))

    singles(deck, compile_dir, actin_block())
    finish(deck, export, 'summary')


def build_zones(compile_dir, panels, export):
    deck = start('zones')
    singles(deck, compile_dir, ZONE_BLOCK)
    V = lambda n: check(panels / 'violins' / n)
    bad.tight_grid(deck, 'Actin by Zone: Normalized by MFI over Full Synapse',
                   [V(f'zones_actin_{TAG}_{t}.png') for t in TPS], TP_LBL, key=True)
    finish(deck, export, 'zones')


def build_threshold(compile_dir, export):
    deck = start('threshold')
    singles(deck, compile_dir, THRESHOLD_BLOCK)
    finish(deck, export, 'threshold')


def montage_png(cond, sub):
    folder = ROOT / cond / 'converted' / 'cropped' / 'split_channels' / 'prog_fixed_cells' / sub
    png = bad.first_page(folder)
    if png is None:
        MISSING.append(folder / 'montages_filtered')
    return png


def montage_slide(deck, title, sides, ppi):
    """sides = [(label, png), ...]: CAT left, FMC right, both at ppi. Same slide as our other
    CART montage decks (insert_CART_CatB_MT_actin_nuc_montages_slides.build_compare_slide):
    black background, white 28 pt title, white 16 pt labels, CELL_W x IMG_H cells. Written
    out here only so a QC slide with one side (no partner page) shows no '(missing)' box."""
    m = bad.mont
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.titles.append(title)
    m.set_slide_background(s, m.BLACK)
    m.add_textbox(s, title, m.TITLE_LEFT, m.TITLE_TOP, m.TITLE_WIDTH, m.TITLE_HEIGHT,
                  font_pt=m.TITLE_FONT_PT, color=m.WHITE, bold=True)
    for (lab, png), (left, top) in zip(sides, m.CELL_POSITIONS):
        if not lab and png is None:
            continue                     # QC slide: this side has no further page
        m.add_textbox(s, lab, left, top, m.CELL_W, m.LABEL_H, font_pt=m.LABEL_FONT_PT, color=m.WHITE, bold=True)
        if png is not None and ppi is not None:
            m.add_image_at_ppi(s, png, ppi, left, top + m.LABEL_H, m.CELL_W, m.IMG_H)
        else:
            m.add_textbox(s, '(missing)', left, top + m.LABEL_H + m.IMG_H / 2 - 0.15, m.CELL_W, 0.3,
                          font_pt=14, color=m.WHITE)


def qc_montage_specs():
    """CAT 5 min QC check (FF 2026-10-18): the cells removed because no actin bottom slice
    (synapse plane) was found, beside kept cells of the same fields of view. Rendered by
    tests/render_undetermined_xz_20261008.m with the pipeline's XZ renderer, so they join the
    'xz' scale group. -> [(title, 'xz', [(label, png), (label, png)]), ...]"""
    qc = ROOT / 'CAT_5' / 'converted' / 'cropped' / 'split_channels' / 'prog_fixed_cells' / 'actin' / 'xz_mip_undetermined_check'
    removed = sorted((qc / 'undetermined' / 'montages').glob('montage_cells_*.png'))
    kept = sorted((qc / 'kept_same_fov' / 'montages').glob('montage_cells_*.png'))
    if not removed or not kept:
        MISSING.append(qc)
    title = f'Actin XZ MIP {DASH} CAT 5 min: Removed by QC vs Kept'
    specs = []
    for k in range(max(len(removed), len(kept))):
        sides = []
        if k < len(removed):
            sides.append((f'Removed: No Synapse Plane Found ({k + 1} of {len(removed)})', removed[k]))
        else:
            sides.append(('', None))
        if k < len(kept):
            sides.append((f'Kept, Same Fields ({k + 1} of {len(kept)})', kept[k]))
        specs.append((title, 'xz', sides))
    return specs


def build_montages(export):
    specs = []                       # (title, scale group, sides)
    for entry in MONTAGES:
        variants = entry if isinstance(entry, list) else [entry]
        for t, tl in zip(TPS, TP_LBL):
            tp = t.replace('min', '')
            pairs = [(tl, f'FMC_{tp}', 'FMC')]
            if t == '10min':        # FMC 10 min was acquired in two folders (pooled at compile)
                pairs.append((f'{tl} (FMC Second Acquisition)', 'FMC_10001', 'FMC (FMC_10001)'))
            for suffix, fmc_cond, fmc_lab in pairs:
                for stem, sub, group in variants:
                    specs.append((f'{stem} {DASH} {suffix}', group,
                                  [('CAT', montage_png(f'CAT_{tp}', sub)),
                                   (fmc_lab, montage_png(fmc_cond, sub))]))

    # CAT 5 min QC check, inserted after the last Actin XZ MIP slide.
    last_actin_xz = max(i for i, sp in enumerate(specs) if sp[0].startswith(f'Actin XZ MIP {DASH}'))
    specs[last_actin_xz + 1:last_actin_xz + 1] = qc_montage_specs()

    # One px/in per scale group, pinned to the largest montage in it (over the established
    # slide's CELL_W x IMG_H image box).
    groups: dict[str, list] = {}
    for _, group, sides in specs:
        groups.setdefault(group, []).extend(png for _, png in sides if png is not None)
    group_ppi = {g: bad.mont.compute_group_ppi(p, bad.mont.CELL_W, bad.mont.IMG_H) for g, p in groups.items() if p}
    for g, v in group_ppi.items():
        print(f'  scale group {g:<28s} {v:7.1f} px/in')

    deck = start('montages')
    for title, group, sides in specs:
        montage_slide(deck, title, sides, group_ppi.get(group))
    finish(deck, export, 'montages')


# --- Overview deck (FF 2026-10-20): format + content of FF's hand-made
# CART_CatB_MT_20260702.pptx (montage slides first, then one compile grid per slide), for
# this experiment, plus GzmB amount (log), GzmB+ fraction, zones. Same panels as the other
# decks here; montages at ONE px/in per view.
# (title stem, prog_fixed_cells sub-folder, scale group); a list = variants on consecutive
# slides for each timepoint. FF 2026-10-20: GzmB XZ MIPs first, the two contrasts alternating
# per timepoint; FMC 10 min = the FMC_10 folder only (no separate FMC_10001 slides here).
OV_MONTAGES = [
    [('GzmB XZ MIP (Fixed Contrast 10–100)', r'GzmB\xz_mip\fixed_con_10_100', 'xz'),
     ('GzmB XZ MIP (Fixed Contrast 10–500)', r'GzmB\xz_mip\fixed_con_10_500', 'xz')],
    ('Actin, Synapse Slice', r'actin\synapse\inner_outer\1slice_combined', 'actin_syn_combined'),   # FF 2026-10-20
    # FF 2026-10-20, replacing the per-cell auto-scaled actin + marker composites: MIP over the
    # bottom 3 planes, actin 3-slice mask outlined in red (tests/render_synapse_3mip_outline_20261008.m)
    [('GzmB 3-Slice MIP, Actin Outline (Fixed Contrast 10–100)',
      r'GzmB\synapse\3slice_actin_outline\fixed_con_10_100', 'synapse'),
     ('GzmB 3-Slice MIP, Actin Outline (Fixed Contrast 10–500)',
      r'GzmB\synapse\3slice_actin_outline\fixed_con_10_500', 'synapse')],
    ('MT 3-Slice MIP, Actin Outline', r'MT\synapse\3slice_actin_outline', 'synapse'),
    ('MT XZ MIP with Centrosome', r'MT\centrosome', 'centrosome'),
]
_G = dict(GZMB_BLOCK) | dict(ZONE_BLOCK) | dict(THRESHOLD_BLOCK)
OV_RESULTS = [(t, _G.get(t)) for t in [  # None: violin-only slide (OV_VIOLINS)
    'Synapse Area',   # actin-based (actin_bottom_mask_area; FF 2026-10-20)
    'Centrosome-Synapse Distance',
    'Centrosome-Synapse Distance: Fraction Polarized',
    'GzmB Total Signal',   # scatter only: its violin twin dropped (FF 2026-10-20 "keep 29")
    f'GzmB Total Signal CDF {DASH} CAT vs FMC',
    'GzmB+ Fraction at Three Total-Signal Cutoffs',   # after the amount + CDF (FF 2026-10-21)
    'GzmB Peak Signal (All Cells, Ungated)',   # scatter, as Total Signal (FF 2026-10-20)
    'GzmB Synapse Total Signal',               # scatter only: violin twin dropped (FF 2026-10-20)
    'GzmB Synapse MFI (3-Slice MIP)',
    f'GzmB {SUB50} ({MU}m from Cell Bottom)',
    f'GzmB {SUB75} ({MU}m from Cell Bottom)',
    f'GzmB {SUB50} (Normalized by Cell Height)',
    f'GzmB {SUB75} (Normalized by Cell Height)',
    f'GzmB {SUB50}: Fraction Polarized',
    f'GzmB {SUB75}: Fraction Polarized',
    f'GzmB Fraction Within 1 {MU}m of Centrosome',
    f'GzmB Fraction Within 2 {MU}m of Centrosome',
    f'GzmB Fraction Within 3 {MU}m of Centrosome',
    # 'Centrosome XY Offset / Synapse Radius: Fraction Centered' and 'Centrosome Synapse Zone':
    # dropped from the overview (FF 2026-10-21)
]]
# Linear metrics as one-axis violins (CAT / FMC at 5, 10, 15 min; FF 2026-10-20: the pMLC
# zone-violin style instead of the timecourse scatter), from
# tests/run_gzmb_20261008_metric_violins.m: title -> (group, field). The log GzmB amount
# scatters, the 3-panel log violins and the fraction bars stay as they are.
OV_VIOLINS = {
    'Synapse Area': ('qc', 'actin_bottom_mask_area'),
    'Centrosome-Synapse Distance': ('qc', 'centrosome_center_z_cell_bottom_distance'),
    # (GzmB total / peak / synapse total: the log scatters, not violins; FF 2026-10-20)
    'GzmB Synapse MFI (3-Slice MIP)': ('qc', 'GzmB_synapse_MFI_3mip'),  # log10 y (FF 2026-10-20)
    f'GzmB {SUB50} ({MU}m from Cell Bottom)': ('gzmb_plus', 'GzmB_z50_cell_bottom_distance'),
    f'GzmB {SUB75} ({MU}m from Cell Bottom)': ('gzmb_plus', 'GzmB_z75_cell_bottom_distance'),
    f'GzmB {SUB50} (Normalized by Cell Height)': ('gzmb_plus', 'GzmB_z50_cell_bottom_distance_norm_cell_height'),
    f'GzmB {SUB75} (Normalized by Cell Height)': ('gzmb_plus', 'GzmB_z75_cell_bottom_distance_norm_cell_height'),
    **{f'GzmB Fraction Within {k} {MU}m of Centrosome': ('gzmb_plus', f'GzmB_frac_around_cent_{k}um') for k in (1, 2, 3)},
}


def ov_violin_png(title):
    group, field = OV_VIOLINS[title]
    d = newest(f'metric_violins_{group}_*')
    if d is None:
        MISSING.append(ROOT / f'metric_violins_{group}_*')
        return None
    return check(d / 'violins' / f'puncta_{field}_{TAG}.png')


# CART_CatB_MT_20260702.pptx geometry (inches): montage slides = black, 24 pt white title,
# 14 pt white CAT / FMC labels and a '5 μm scalebar' key; results = 28 pt black title.
OV_LEFTS, OV_CELL_W = (0.05, 6.68), 6.6
OV_IMG_TOP = 0.80
OV_IMG_H = bpd.SLIDE_H - OV_IMG_TOP - 0.05


def ov_montage_slide(deck, title, sides, ppi):
    m = bad.mont
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.titles.append(title)
    m.set_slide_background(s, m.BLACK)
    m.add_textbox(s, title, 0.05, 0.05, bpd.SLIDE_W - 0.1, 0.45, font_pt=24, color=m.WHITE, bold=True)
    # Key in the label row, left of 'CAT' (FF's deck had it top-left; long titles here reach it).
    m.add_textbox(s, f'5 {MU}m scalebar', 0.05, 0.5, 1.91, 0.3, font_pt=14, color=m.WHITE, bold=True)
    for (lab, png), left in zip(sides, OV_LEFTS):
        if not lab and png is None:
            continue
        m.add_textbox(s, lab, left, 0.5, OV_CELL_W, 0.3, font_pt=14, color=m.WHITE, bold=True)
        if png is not None and ppi is not None:
            m.add_image_at_ppi(s, png, ppi, left, OV_IMG_TOP, OV_CELL_W, OV_IMG_H)
        else:
            m.add_textbox(s, '(missing)', left, OV_IMG_TOP + OV_IMG_H / 2 - 0.15, OV_CELL_W, 0.3,
                          font_pt=14, color=m.WHITE)


class OvDeck(bpd.Deck):
    """bpd.Deck with the 28 pt results-slide title of FF's deck (also used by bad.tight_grid)."""
    def slide(self, title, subtitle=''):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bpd.add_text(s, 0.1, 0.05, bpd.SLIDE_W - 0.2, 0.6, title, 28, bold=True, align=PP_ALIGN.CENTER)
        self.titles.append(title)
        return s


# Every slide states its plane / region and its cells (FF 2026-10-20: "is this based on 1 plane
# or 3?? this sort of thing needs to be specified"). Planes, from the code:
#   synapse slice = the single bottom plane (CART_fixed_cells bot_slice); 3-slice MIP = max over
#   it and the 2 planes above; whole cell = all planes bottom..top (process_GzmB_channel peak /
#   total). Since the 2026-10-20 fix (process_GzmB_channel), GzmB synapse MFI / total are
#   taken inside the actin footprint of their own plane: actin_bottom_mask (synapse slice) /
#   actin_3slice_mask (3-slice MIP), no longer the whole-cell actin MIP outline (actin_MIP_mask).
OV_SHOW = {   # OV_RESULTS key -> slide title
    'Synapse Area': 'Synapse Area (Actin, Synapse Slice)',
    'GzmB Total Signal': 'GzmB Total Signal (Whole Cell)',
    f'GzmB Total Signal CDF {DASH} CAT vs FMC': f'GzmB Total Signal (Whole Cell) CDF {DASH} CAT vs FMC',
    'GzmB Peak Signal (All Cells, Ungated)': 'GzmB Peak Signal (Whole Cell)',
    'GzmB Synapse Total Signal': 'GzmB Synapse Total Signal (3-Slice MIP)',
}
UNGATED = 'All cells passing QC; no GzmB-based gating.'
SLICE_REGION = 'Region: actin footprint on the synapse slice.'
MIP3_REGION = 'Region: actin footprint on the 3-slice MIP.'
AXIAL_REGION = 'Region: actin MIP outline, bottom to top slice.'
CENT_REGION = 'Region: actin MIP outline, bottom to top slice; 3D distance from the centrosome.'
GZMB_PLUS_NOTE = 'GzmB+ = whole-cell GzmB total signal > 10⁶.'
# Italic cell-selection line on slides whose gating is not in the title (FF 2026-10-20).
OV_NOTES = {
    'Synapse Area': 'All cells passing QC.',
    # centrosome_center_MT: argmax of the smoothed MT stack inside the actin MIP outline, every cell
    'Centrosome-Synapse Distance': ('All cells passing QC. Centrosome = brightest point of the '
                                    'smoothed MT stack within the actin MIP outline.'),
    # axial_signal_binning_z_50_75_90(I_GzmB, actin_MIP_mask, bot, top, ...)
    **{f'GzmB {z} ({MU}m from Cell Bottom)': AXIAL_REGION for z in (SUB50, SUB75)},
    **{f'GzmB {z} (Normalized by Cell Height)': AXIAL_REGION for z in (SUB50, SUB75)},
    # prot_cent_assoc_3D: slices bot..top, numerator and denominator (FF 2026-10-21)
    **{f'GzmB Fraction Within {k} {MU}m of Centrosome': CENT_REGION for k in (1, 2, 3)},
    'GzmB Total Signal': UNGATED,
    f'GzmB Total Signal CDF {DASH} CAT vs FMC': UNGATED,
    'GzmB Peak Signal (All Cells, Ungated)': UNGATED,
    'GzmB Synapse Total Signal': f'{MIP3_REGION} {UNGATED}',
    'GzmB Synapse MFI (3-Slice MIP)': f'{MIP3_REGION} {UNGATED}',
}

# Centrosome zone slide (stacked bars): what is counted and the test.
ZONE_NOTE = ("Zone of the centrosome's XY position on the synapse-slice actin footprint. "
             'Stars: chi-square test of the zone distribution, Benjamini-Hochberg corrected.')

# Method line on every fraction-of-cells slide (FF 2026-10-20: '95% and method ... in italics').
FRACTION_NOTE = ('Error bars: 95% Clopper-Pearson binomial confidence intervals. '
                 "Stars: Fisher's exact test, Benjamini-Hochberg corrected.")


def ov_result_slide(deck, title, png, note=''):
    s = deck.slide(title)
    if title.endswith('GzmB+ Cells Only'):
        note = f'{GZMB_PLUS_NOTE} {note}'.strip()
    note_h = (0.4 if len(note) <= 125 else 0.7) if note else 0   # 16 pt: ~125 characters per line
    box_h = bpd.SLIDE_H - 0.85 - note_h
    bpd.fit_in_box(s, png, 0.1, 0.75, bpd.SLIDE_W - 0.2, box_h)
    if note:
        tb = bpd.add_text(s, 0.1, bpd.SLIDE_H - 0.1 - note_h, bpd.SLIDE_W - 0.2, note_h, note, 16,
                          align=PP_ALIGN.CENTER)
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.italic = True


# Fraction-of-cells: ONE wide figure per cutoff, one slide each (FF 2026-10-20: the 3-tile grid's
# text was far too small). field -> (comparison, unit, title stem); the cutoffs come from the
# compile's polarization\<experiment>\<field>_pol_thresh<t>.png files (plotPolarizationStacked).
POL_SLIDES = {
    # '≤': plotPolarizationStacked counts value <= cutoff (FF 2026-10-20); GzmB+ stays '>'.
    'centrosome_center_z_cell_bottom_distance': ('≤', f' {MU}m', 'Centrosome-Synapse Distance: Fraction Polarized'),
    'GzmB_total_sig': ('>', '', 'GzmB+ Fraction (Whole-Cell Total Signal'),
    'GzmB_z50_cell_bottom_distance': ('≤', f' {MU}m', f'GzmB {SUB50}: Fraction Polarized'),
    'GzmB_z75_cell_bottom_distance': ('≤', f' {MU}m', f'GzmB {SUB75}: Fraction Polarized'),
    'centrosome_xy_offset_norm': ('≤', ' × Synapse Radius', 'Centrosome XY Offset: Fraction Centered'),
}
SUP = str.maketrans('0123456789.', '⁰¹²³⁴⁵⁶⁷⁸⁹·')


def pol_cutoff_files(compile_dir, field):
    """[(cutoff, png)] sorted by cutoff; filename token = MATLAB %g with '.' -> 'p'."""
    out = []
    for p in glob.glob(str(compile_dir / 'polarization' / '*' / f'{field}_pol_thresh*.png')):
        tok = Path(p).stem.split('_pol_thresh')[1].replace('p', '.')
        out.append((float(tok), Path(p)))
    return sorted(out)


def ov_fraction_slides(deck, compile_dir, field):
    op, unit, stem = POL_SLIDES[field]
    files = pol_cutoff_files(compile_dir, field)
    if not files:
        MISSING.append(compile_dir / 'polarization' / f'{field}_pol_thresh*.png')
    for t, png in files:
        if field == 'GzmB_total_sig':      # 10^5, 10^5.5, 10^6
            val = '10' + f'{math.log10(t):.1f}'.rstrip('0').rstrip('.').translate(SUP)
            title = f'{stem} {op} {val})'
        else:
            title = f'{stem} ({op} {t:g}{unit})'
        note = FRACTION_NOTE
        if field.startswith('GzmB_z'):
            title += f' {DASH} GzmB+ Cells Only'      # ov_result_slide adds the GzmB+ definition
        elif field == 'centrosome_xy_offset_norm':
            note = f'Synapse radius = √(synapse-slice actin area / π). {note}'
        ov_result_slide(deck, title, png, note)


def ov_actin_zone_slide(deck):
    """Actin MFI by zone (FF 2026-10-20), the actin panels' 3-zone violins
    (run_gzmb_20261008_extended_panels, compile QC gates)."""
    d = newest('actin_extended_panels_*')
    if d is None:
        MISSING.append(ROOT / 'actin_extended_panels_*')
        return
    bad.tight_grid(deck, 'Actin by Zone, Synapse Slice (MFI / Synapse Mean)',
                   [check(d / 'violins' / f'zones_actin_{TAG}_{t}.png') for t in TPS], TP_LBL, key=True)


def ov_profile_slides(deck, marker):
    """pMLC-style profile panels (tests/run_gzmb_20261008_extended_panels.m, marker mode):
    actin on the synapse slice, the marker on the 3-slice synapse MIP; GzmB panels = GzmB+
    cells only (the compile gate), MT panels = every QC cell."""
    d = newest(f'{marker.lower()}_extended_panels_*')
    if d is None:
        MISSING.append(ROOT / f'{marker.lower()}_extended_panels_*')
        return
    P = lambda n: check(d / 'profiles' / n)  # noqa: E731
    pair = f'Actin (Synapse Slice) and {marker} (3-Slice MIP)'
    gate = f' {DASH} GzmB+ Cells' if marker == 'GzmB' else ''
    # Zone MFIs: the 3-zone violins on the 3-SLICE MIP with zones cut on the actin 3-slice
    # mask (synapse_zone_position _3mip fields, cfg.zone_suffix), the plane + mask of the
    # profiles (FF 2026-10-21; they were on the synapse slice before). GzmB: cells gated on
    # synapse GzmB, not whole-cell GzmB+ (FF 2026-10-20), from the 'GzmB_zones' run.
    zd, ztitle = d, f'{marker} by Zone, 3-Slice MIP (MFI / Synapse Mean){gate}'
    if marker == 'GzmB':
        zd = newest('gzmb_zones_mfi3mip_extended_panels_*')
        if zd is None:
            MISSING.append(ROOT / 'gzmb_zones_mfi3mip_extended_panels_*')
            zd = d
        ztitle = f'{marker} by Zone, 3-Slice MIP (MFI / Synapse Mean)'
    bad.tight_grid(deck, ztitle,   # short: CAT / FMC key top right
                   [check(zd / 'violins' / f'zones_{marker.lower()}_{TAG}_{t}.png') for t in TPS], TP_LBL, key=True)
    if marker == 'GzmB':   # threshold = params.gzmb_synapse_mfi_thresh_3mip (compile config)
        tb = bpd.add_text(deck.prs.slides[-1], 0.1, bpd.SLIDE_H - 0.5, bpd.SLIDE_W - 0.2, 0.4,
                          'Cells: GzmB synapse MFI > 10 on the 3-slice synapse MIP '
                          '(not the whole-cell GzmB+ gate).', 16, align=PP_ALIGN.CENTER)
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.italic = True
    bad.tight_grid(deck, f'{pair} Profiles{gate}',
                   [P(f'both_catfmc_profile_{TAG}_{t}.png') for t in TPS] +
                   [P(f'both_catfmc_grad_{TAG}_{t}.png') for t in TPS], TP_LBL)
    bad.tight_grid(deck, f'FMC {MINUS} CAT: {pair}{gate}',
                   [P(f'diff_both_profile_{t}.png') for t in TPS] +
                   [P(f'diff_both_grad_{t}.png') for t in TPS], TP_LBL)
    # tc_* panels draw the marker only (plot_pmlc_extended_panels: n_ch = 'mark'), no actin.
    bad.tight_grid(deck, f'{marker} (3-Slice MIP) Profile Time Course{gate}',
                   [P(f'tc_{ct}_{kind}_{TAG}.png') for kind in ('profile', 'grad') for ct in ('CAT', 'FMC')],
                   ['CAT', 'FMC'])   # 2 x 2 (profile row, gradient row): four in a row were tiny


def build_overview(compile_dir, export):
    bpd.OUT = OUT_DIR / f'CART_GzmB_MT_20261008_{TODAY}.pptx'
    deck = OvDeck()
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.titles.append('Title')
    bpd.add_text(s, 0.83, 2.4, 11.67, 1.2, 'CAR T Cell Project: Granzyme B Polarization', 40,
                 bold=True, align=PP_ALIGN.CENTER)
    bpd.add_text(s, 0.83, 3.75, 11.67, 0.5, f'Oct 8, 2026 Experiment {DASH} CAT vs FMC at 5, 10, 15 min', 24,
                 align=PP_ALIGN.CENTER)
    today = datetime.date.today()
    bpd.add_text(s, 4.37, 4.6, 4.6, 0.5, f'{today:%B} {today.day}, {today.year}', 20, align=PP_ALIGN.CENTER)

    # Montages: per view, every timepoint; variants of a view alternate within a timepoint.
    specs = []
    for entry in OV_MONTAGES:
        for t, tl in zip(TPS, TP_LBL):
            tp = t.replace('min', '')
            for stem, sub, group in (entry if isinstance(entry, list) else [entry]):
                specs.append((f'{stem} {DASH} {tl}', group,
                              [('CAT', montage_png(f'CAT_{tp}', sub)), ('FMC', montage_png(f'FMC_{tp}', sub))]))
    # QC removed-vs-kept slides live in the montages deck only (FF 2026-10-20: drop from overview)
    groups: dict[str, list] = {}
    for _, group, sides in specs:
        groups.setdefault(group, []).extend(png for _, png in sides if png is not None)
    group_ppi = {g: bad.mont.compute_group_ppi(p, OV_CELL_W, OV_IMG_H) for g, p in groups.items() if p}
    for g, v in group_ppi.items():
        print(f'  scale group {g:<28s} {v:7.1f} px/in')
    for title, group, sides in specs:
        ov_montage_slide(deck, title, sides, group_ppi.get(group))

    for title, rel in OV_RESULTS:   # (no Cells per Condition slide: FF 2026-10-20)
        if title == 'Centrosome Synapse Zone':
            zp = sorted(glob.glob(str(compile_dir / 'zone_fractions' / '*' / 'centrosome_synapse_zone.png')))
            if not zp:
                MISSING.append(compile_dir / 'zone_fractions' / 'centrosome_synapse_zone.png')
                continue
            ov_result_slide(deck, title, Path(zp[0]), ZONE_NOTE)
            continue
        if rel is None:   # violin-only slide: its field decides the GzmB+ title suffix
            rel = f'{OV_VIOLINS[title][1]}_grid.png'
        if rel.endswith('_polarization_grid.png'):
            ov_fraction_slides(deck, compile_dir, rel[:-len('_polarization_grid.png')])
            continue
        png = ov_violin_png(title) if title in OV_VIOLINS else check(compile_dir / 'grid_panels' / rel)
        ov_result_slide(deck, gated_title(OV_SHOW.get(title, title), rel), png, OV_NOTES.get(title, ''))
        if title == f'GzmB Fraction Within 3 {MU}m of Centrosome':   # then zones (actin, GzmB) + profiles, MT
            ov_actin_zone_slide(deck)
            ov_profile_slides(deck, 'GzmB')
            ov_profile_slides(deck, 'MT')
    finish(deck, export, 'overview')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--deck', default='all', choices=['summary', 'zones', 'threshold', 'montages', 'overview', 'all'])
    ap.add_argument('--compile', help='compile dir (default: newest <root>\\compiled_2*)')
    ap.add_argument('--export', help='also export one PNG per slide (named by title) under this dir')
    args = ap.parse_args()
    compile_dir = Path(args.compile) if args.compile else newest('compiled_2*')
    panels = newest('actin_extended_panels_*')
    if compile_dir is None:
        print('[error] no compile found; run run/run_gzmb_20261008.m')
        return 1
    print(f'compile: {compile_dir}\npanels:  {panels}')
    if panels is None and args.deck in ('summary', 'zones', 'all'):
        print('[error] no actin_extended_panels_*; run tests/run_gzmb_20261008_extended_panels.m')
        return 1
    if args.deck in ('summary', 'all'):
        build_summary(compile_dir, panels, args.export)
    if args.deck in ('zones', 'all'):
        build_zones(compile_dir, panels, args.export)
    if args.deck in ('threshold', 'all'):
        build_threshold(compile_dir, args.export)
    if args.deck in ('montages', 'all'):
        build_montages(args.export)
    if args.deck in ('overview', 'all'):
        build_overview(compile_dir, args.export)
    if MISSING:
        print(f'\n[MISSING] {len(MISSING)} panel(s):')
        for m in MISSING:
            print(f'  {m}')
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
