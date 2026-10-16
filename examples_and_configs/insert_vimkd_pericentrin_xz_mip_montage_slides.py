"""
insert_vimkd_pericentrin_xz_mip_montage_slides.py

siCtrl vs siVim X-Z MIP montage deck (actin magenta + pericentrin yellow) for the fixed Jurkat
VimentinKD pericentrin experiments, PER EXPERIMENT (fixed cells are never pooled): one slide per
experiment with montage pages 1 and 2 stacked (5 x 5 grids: the first 50 cells in the analysis),
siCtrl left, siVim right. A last slide shows the vimentin-stained Apr 6 coverslips (transfection2_48h, vimentin
channel only, ONE display range for every cell so the knockdown shows; the range is read from the run log and
written on the slide). Montages from FIJI_macros\\run_vimkd_pericentrin_xz_mip_montages.ps1 (Aashli's
create_xz_mip_montage_all_cells recipe; 5 um bar embedded).

Same style and machinery as insert_vimkd_pericentrin_phys_scale_slides.py (imported, not copied):
16:9 black slides, white bold 28 pt title, white bold 16 pt condition labels, siCtrl left / siVim
right. PNGs are inserted at one pixels-per-inch PER EXPERIMENT, so
siCtrl and siVim (and both pages) of an experiment share a scale; the embedded 5 um bar on every
montage gives the scale, and its length differs between experiments (FF 2026-10-05: ok).

Usage:
    conda run -n PPT_editing python examples_and_configs/insert_vimkd_pericentrin_xz_mip_montage_slides.py
"""

import os
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import insert_vimkd_pericentrin_phys_scale_slides as PS

RUN = Path("M:/FF/FF_4TB_2_Backup_fullHD/Vimentin_Project_2ndharddrive/VimentinKD_NucleusData_Fixed/"
           "results_compilation/VimKD_pericentrin_xz_mip_montages_nolabels_20261005")   # FF 2026-10-05: 5 x 5 pages, no "Cell N"
OUTPUT_PATH = Path("K:/FF/PPT/PPT_autogeneration/Fixed Jurkats, Miscellaneous/VimentinKD/"
                   "VimKD_pericentrin_xz_mip_montages_20261005.pptx")
EXPERIMENTS = [("apr6", "Apr 6, 2022"), ("may20", "May 20, 2022"), ("jun28", "Jun 28, 2022")]
CONDS = [("siCtrl", "siCtrl"), ("siVim", "siVim")]       # FF 2026-10-05: no "(KD)"
PAGES = [1, 2]
TITLE = "{date}: Actin + pericentrin, X-Z MIPs"          # FF 2026-10-05: "Actin", not "F-actin"
# channel legend, upper left of every slide: the montages' Fiji LUTs (Magenta, Yellow)
LEGEND = [("Actin", RGBColor(0xFF, 0x00, 0xFF)), ("Pericentrin", RGBColor(0xFF, 0xFF, 0x00))]
LEGEND_PT = 20
# vimentin knockdown slide (FIJI_macros\run_vimkd_vimentin_xz_mip_montages.ps1, single channel, fixed range)
VIM_RUN = Path("M:/FF/FF_4TB_2_Backup_fullHD/Vimentin_Project_2ndharddrive/VimentinKD_NucleusData_Fixed/"
               "results_compilation/VimKD_vimentin_xz_mip_montages_cyan_20261005")   # FF 2026-10-05: cyan, range 110-10000
VIM_LOG = VIM_RUN / "run_vimentin_20261005.log"
VIM_TAG = "apr6vim"
VIM_TITLE = "Apr 6, 2022 (different cells): Vimentin, X-Z MIPs"   # FF 2026-10-05
VIM_LEGEND = [("Vimentin", RGBColor(0x00, 0xFF, 0xFF))]

# One slide per experiment (FF 2026-10-05 "use space well"): the 5 x 5 pages are ~2.5x wider than tall,
# so one page per half-slide filled only a third of the height; pages 1 and 2 are stacked instead.
#   siCtrl p1 | siVim p1
#   siCtrl p2 | siVim p2
#   each montage has its own label above it: "siCtrl (set 1)" ... (FF 2026-10-05)
COL_GAP = 0.10
CELL_W = (PS.SLIDE_W - 2 * PS.GRID_LEFT - COL_GAP) / 2      # 6.517"
ROW_GAP = 0.05
BOTTOM = PS.SLIDE_H - 0.05
CELL_H = (BOTTOM - PS.GRID_TOP - 2 * PS.LABEL_H - ROW_GAP) / 2   # 3.10" per montage


def montage(tag: str, cond: str, page: int, run: Path = RUN) -> Path:
    return run / f"{tag}_{cond}" / f"{tag}_{cond}_xz_mip_montage_{page:03d}.png"


def add_legend(slide, legend=LEGEND) -> None:
    """Channel names in their colors, bold, in the upper-left corner of the title band."""
    box = slide.shapes.add_textbox(Inches(PS.TITLE_LEFT), Inches(PS.TITLE_TOP), Inches(2.6), Inches(PS.TITLE_HEIGHT))
    tf = box.text_frame
    tf.margin_left, tf.margin_top = Inches(0.05), Inches(0.02)
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    for i, (name, rgb) in enumerate(legend):
        run = para.add_run()
        run.text = name if i == 0 else "   " + name
        run.font.size, run.font.bold = Pt(LEGEND_PT), True
        run.font.color.rgb = rgb


def vim_display_range() -> tuple:
    """The single display range of the vimentin run, from its log ('  cellN ch1 display lo-hi'); asserts one range."""
    txt = VIM_LOG.read_text(encoding="utf-8", errors="replace")
    ranges = set(re.findall(r"cell\d+ ch1 display (\S+)-(\S+)", txt))
    assert len(ranges) == 1, f"vimentin run has {len(ranges)} display ranges, expected one: {sorted(ranges)[:5]}"
    lo, hi = ranges.pop()
    return float(lo), float(hi)


def add_montage_slide(prs, title: str, legend, grid, note: str = "") -> None:
    """One slide: grid[page][cond] montages, siCtrl left / siVim right, pages stacked, one PPI for the slide."""
    missing = [str(p) for row in grid for p in row if not PS._exists_long(p)]
    assert not missing, f"missing montage(s): {missing}"
    # one PPI per slide (FF 2026-10-05: bars may differ between experiments): the smallest at which
    # all four montages fit their cells, so this slide is as large as the page allows
    ppi = PS.compute_slide_ppi([p for row in grid for p in row], CELL_W, CELL_H)
    print(f"{title}: PPI {ppi:.1f}, 5 um bar = {5 / 0.065 / ppi:.3f} in")   # Fiji Scale Bar width=5 um at 0.065 um/px

    s = prs.slides.add_slide(prs.slide_layouts[6])
    PS.set_slide_background(s, PS.BLACK)
    PS.add_textbox(s, title, PS.TITLE_LEFT, PS.TITLE_TOP, PS.TITLE_WIDTH, PS.TITLE_HEIGHT,
                   font_pt=PS.TITLE_FONT_PT, color=PS.WHITE, bold=True)
    add_legend(s, legend)
    if note:   # upper left, under the legend (FF 2026-10-05)
        box = PS.add_textbox(s, note, PS.TITLE_LEFT, PS.TITLE_TOP + PS.TITLE_HEIGHT - 0.05, 4.0, 0.32,
                             font_pt=LEGEND_PT - 6, color=PS.WHITE, bold=True)
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    # rows packed: label + montage per row, each row as tall as its tallest montage, ROW_GAP between,
    # the block centered under the title (fixed half-height cells left black bands)
    row_h = [max(PS._png_dims(p)[1] for p in row) / ppi for row in grid]
    block = sum(row_h) + len(grid) * PS.LABEL_H + ROW_GAP * (len(grid) - 1)
    tops = [PS.GRID_TOP + (BOTTOM - PS.GRID_TOP - block) / 2]   # label tops
    for h in row_h[:-1]:
        tops.append(tops[-1] + PS.LABEL_H + h + ROW_GAP)
    for k, (_, label) in enumerate(CONDS):
        left = PS.GRID_LEFT + k * (CELL_W + COL_GAP)
        for r, (page, row) in enumerate(zip(PAGES, grid)):
            PS.add_textbox(s, f"{label} (set {page})", left, tops[r], CELL_W, PS.LABEL_H,
                           font_pt=PS.LABEL_FONT_PT, color=PS.WHITE, bold=True)
            PS.add_image_at_ppi(s, row[k], ppi, left, tops[r] + PS.LABEL_H, CELL_W, row_h[r])
    print(f"  {', '.join(p.name for row in grid for p in row)}")


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PS.SLIDE_W), Inches(PS.SLIDE_H)
    for tag, date in EXPERIMENTS:
        add_montage_slide(prs, TITLE.format(date=date), LEGEND,
                          [[montage(tag, c, page) for c, _ in CONDS] for page in PAGES])
    if VIM_LOG.exists():
        lo, hi = vim_display_range()
        add_montage_slide(prs, VIM_TITLE, VIM_LEGEND,
                          [[montage(VIM_TAG, c, page, VIM_RUN) for c, _ in CONDS] for page in PAGES],
                          note=f"Display range {lo:.0f}–{hi:.0f}, same for all cells")
    else:
        print(f"no vimentin run log ({VIM_LOG}): vimentin slide skipped")

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))
    print(f"Done. {len(prs.slides)} slides -> {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
