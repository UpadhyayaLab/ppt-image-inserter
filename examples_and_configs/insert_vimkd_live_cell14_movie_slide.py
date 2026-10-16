"""
insert_vimkd_live_cell14_movie_slide.py

Independent one-slide deck (FF 2026-10-04): the existing XZ-MIP movie of live siCtrl
cell 14 (20220519; centrin + nucleus, Fiji export from 2022) copied as is -- FF decided
not to remake it. Source = the movie embedded in Aashli's deck:
  J:/FF/vim_live_cell/vimentinKD_3D_timelapse_centrin_nucleus/XZ_MIPs_all.pptx
  slide 14, shape "XZ-MIP_cell14" -> ppt/media/media89.mp4, poster ppt/media/image89.png
(read from the .pptx zip; the deck itself is not opened or modified).

Style = insert_vimkd_centrosome_slide.py (16:9, 48 pt bold title, no footer).

Usage:
    conda run -n PPT_editing python examples_and_configs/insert_vimkd_live_cell14_movie_slide.py
"""

import os
import sys
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import insert_channel_misalignment_slides as CM

SRC_DECK = Path("J:/FF/vim_live_cell/vimentinKD_3D_timelapse_centrin_nucleus/XZ_MIPs_all.pptx")
MOVIE, POSTER = "ppt/media/media89.mp4", "ppt/media/image89.png"   # shape "XZ-MIP_cell14", slide 14

OUT_DIR = Path("K:/FF/PPT/PPT_autogeneration/Fixed Jurkats, Miscellaneous/VimentinKD")
OUT = OUT_DIR / "VimKD_live_siCtrl_cell14_XZ_MIP_20261004.pptx"
TITLE = "siCtrl cell 14 (May 19, 2022): centrosome does not polarize"
TITLE_PT, TITLE_H = 48, 1.65


def main():
    src = OUT_DIR / "source_images" / "cell14"
    src.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(SRC_DECK) as z:
        for member in (MOVIE, POSTER):
            (src / Path(member).name).write_bytes(z.read(member))
    movie, poster = src / Path(MOVIE).name, src / Path(POSTER).name
    with Image.open(poster) as im:
        aspect = im.width / im.height

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(CM.W), Inches(CM.H)
    s = CM.M._new_slide(prs)
    CM.M.add_textbox(s, TITLE, CM.MARGIN, CM.TITLE_TOP, CM.IMG_W, TITLE_H,
                     font_pt=TITLE_PT, color=CM.BLACK, bold=True)
    s.shapes[-1].text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    top = CM.TITLE_TOP + TITLE_H + 0.08
    box_w, box_h = CM.IMG_W, CM.H - top - CM.MARGIN
    w = min(box_w, box_h * aspect)
    h = w / aspect
    left = CM.MARGIN + (box_w - w) / 2
    s.shapes.add_movie(str(movie), Inches(left), Inches(top), Inches(w), Inches(h),
                       poster_frame_image=str(poster), mime_type="video/mp4")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if OUT.exists():
        CM.M.backup_presentation(str(OUT), backup_base=str(OUT_DIR / "backups"))
    prs.save(str(OUT))
    print(f"movie {movie.stat().st_size} B, aspect {aspect:.2f} -> {OUT}")


if __name__ == "__main__":
    main()
