"""
insert_MT_near_invag_siCtrl_vs_DMSO_slides.py

Build the "microtubule (MT) density near nuclear invaginations" deck comparing
two control arms across two experiments:

  - siCtrl        -- VimKD experiment, 20220523 - Vim KD bTub AcTub (n=107 filtered)
  - DMSO baseline -- Nocodazole experiment, prog_fixed_Noco_MT_20240123 (n=77 filtered)

Both experiments' results are merged into a single geom-density comparison folder
(VimKD_siCtrl_vs_NocoDMSO_geomdensity_comparison_YYYYMMDD) whose MT_loc_wrto_invag/
subfolder holds a curated, report-ordered set of comparison figures. Each PNG
already overlays BOTH conditions side by side (paired violins / Mean±SEM
profiles), so this deck is a section-organized insert of pre-made figures, not a
re-plot.

MT is a cytoplasmic (struct_out_nuc) channel: the question is whether MT density
enriches in / tracks nuclear invaginations and concave NE. Hybrid layout -- wide
Mean±SEM density profiles one-per-slide, violin sets grouped into per-section
grids. Self-contained blank 16:9 deck. Missing panels are skipped; a previous
deck at the output path is backed up first.

Usage:
    conda run -n PPT_editing python examples_and_configs/insert_MT_near_invag_siCtrl_vs_DMSO_slides.py
    conda run -n PPT_editing python examples_and_configs/insert_MT_near_invag_siCtrl_vs_DMSO_slides.py --list
"""

import math
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from ppt_image_inserter import (  # noqa: E402
    backup_presentation,
    safe_path,
    resolve_latest_compile,
    compile_date_from,
)

# ---------------------------------------------------------------------------
# Paths -- resolve the newest matching compile so the deck tracks recompiles.
# ---------------------------------------------------------------------------
RESULTS = Path(
    "J:/FF/vim_data/from_AP_04122024/20220523 - Vim KD bTub AcTub/results_compilation"
)
ROOT = resolve_latest_compile(
    RESULTS, "VimKD_siCtrl_vs_NocoDMSO_geomdensity_comparison_*"
)
CURATED = ROOT / "MT_loc_wrto_invag"     # 02a-c avg, 03a-e enrichment, 04a-c corr, 06 orient, 07 nuc shape, 09 counts
COMPILED = compile_date_from(ROOT)        # -> "2026-08-26"

OUTPUT_PATH = Path(
    "K:/FF/PPT/PPT_autogeneration/Fixed Jurkats, Miscellaneous/"
    "nuc_mesh_struct_outside_nuc/MT density near nuclear invaginations, siCtrl vs DMSO.pptx"
)

DECK_TITLE = "Microtubule density near nuclear invaginations"
# n's from cell_counts.csv (FilteredCells column); re-check on rebuild.
DECK_SUBTITLE = (
    "Fixed Jurkats · cross-experiment comparison · "
    "siCtrl (VimKD 20220523, n=107) vs DMSO baseline (Noco 20240123, n=77) · "
    "MT = cytoplasmic (struct_out_nuc), sampled near the NE · compiled " + COMPILED
)

# Shared caption fragments for the paired-condition violins.
PAIR_ENR = "siCtrl vs DMSO  ·  one dot per cell  ·  ref line 1 (enriched > 1)"
PAIR_CORR = "siCtrl vs DMSO  ·  one dot per cell  ·  ref line 0"

# ---------------------------------------------------------------------------
# Colors / layout (16:9)  -- shared with the SUN1 / MT geom-density decks.
# ---------------------------------------------------------------------------
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
FIELD_COLOR = RGBColor(0x2E, 0x5A, 0x88)   # full metric/field name (tracking ID)
DIVIDER_BG = RGBColor(0xF0, 0xF0, 0xF0)

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.10
GAP = 0.16

TITLE_LEFT = MARGIN
TITLE_TOP = 0.05
TITLE_WIDTH = SLIDE_W - 2 * MARGIN
TITLE_HEIGHT = 0.55
TITLE_FONT_PT = 28

IMG_TOP = 0.66
IMG_BOX_W = SLIDE_W - 2 * MARGIN
IMG_BOX_H = 6.36

FOOTER_TOP = 7.06
FOOTER_WIDTH = SLIDE_W - 2 * MARGIN
FOOTER_HEIGHT = 0.40
FOOTER_FONT_PT = 9

CAPTION_HEIGHT = 0.55


# ---------------------------------------------------------------------------
# Slide helpers  (same layout engine as insert_MT_geomdensity_DMSO_slides.py)
# ---------------------------------------------------------------------------
def add_textbox(slide, text, left, top, width, height, font_pt, color,
                bold=False, italic=False, align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.text = text
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.runs[0]
    run.font.size = Pt(font_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def title_font_for(text):
    n = len(text)
    if n <= 52:
        return TITLE_FONT_PT
    if n <= 70:
        return 24
    if n <= 90:
        return 20
    return 18


def add_image_in_box(slide, image_path, box_left, box_top, box_w, box_h):
    """Add an image scaled to fit within (box_w, box_h), centered in the box."""
    pic = slide.shapes.add_picture(
        safe_path(str(image_path)), Inches(box_left), Inches(box_top),
        width=Inches(box_w))
    actual_h_in = pic.height / 914400.0
    if actual_h_in > box_h:
        sp = pic._element
        sp.getparent().remove(sp)
        pic = slide.shapes.add_picture(
            safe_path(str(image_path)), Inches(box_left), Inches(box_top),
            height=Inches(box_h))
        actual_w_in = pic.width / 914400.0
        pic.left = Inches(box_left + (box_w - actual_w_in) / 2)
    else:
        pic.top = Inches(box_top + (box_h - actual_h_in) / 2)
    return pic


def set_slide_background(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def _new_slide(prs, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, bg)
    return slide


def build_title_slide(prs, title, subtitle):
    slide = _new_slide(prs)
    add_textbox(slide, title, MARGIN, 2.7, SLIDE_W - 2 * MARGIN, 1.3,
                font_pt=38, color=BLACK, bold=True)
    add_textbox(slide, subtitle, MARGIN, 4.1, SLIDE_W - 2 * MARGIN, 1.1,
                font_pt=16, color=GREY, italic=True)


def build_divider_slide(prs, title, subtitle=""):
    slide = _new_slide(prs, bg=DIVIDER_BG)
    add_textbox(slide, title, MARGIN, 3.0, SLIDE_W - 2 * MARGIN, 1.2,
                font_pt=40, color=BLACK, bold=True)
    if subtitle:
        add_textbox(slide, subtitle, MARGIN, 4.25, SLIDE_W - 2 * MARGIN, 0.9,
                    font_pt=18, color=GREY, italic=True)


def build_single_slide(prs, title, image_path, caption, footer=None):
    slide = _new_slide(prs)
    add_textbox(slide, title, TITLE_LEFT, TITLE_TOP, TITLE_WIDTH, TITLE_HEIGHT,
                font_pt=title_font_for(title), color=BLACK, bold=True)
    cap_h = CAPTION_HEIGHT if caption else 0.0
    img_h = IMG_BOX_H - cap_h
    add_image_in_box(slide, str(image_path), MARGIN, IMG_TOP, IMG_BOX_W, img_h)
    if caption:
        add_textbox(slide, caption, MARGIN, IMG_TOP + img_h, IMG_BOX_W, cap_h,
                    font_pt=12, color=GREY)
    add_textbox(slide, footer or rel_footer(image_path), MARGIN, FOOTER_TOP,
                FOOTER_WIDTH, FOOTER_HEIGHT, font_pt=FOOTER_FONT_PT, color=GREY)
    return slide


def add_labeled_caption(slide, lines, left, top, width, height):
    """Multi-line caption; each line is (text, font_pt, color). Used to show a
    human hint plus the FULL filename stem (the tracking ID) under each tile."""
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    for i, (text, pt, color) in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = text
        run.font.size = Pt(pt)
        run.font.color.rgb = color
    return box


def _tile_caption(slide, cap, path, left, top, width, height, cap_pt):
    lines = []
    if cap:
        lines.append((cap, min(cap_pt, 10), GREY))
    lines.append((Path(str(path)).stem, 8, FIELD_COLOR))   # full field name for tracking
    add_labeled_caption(slide, lines, left, top, width, height)


def build_grid_slide(prs, title, entries, footer_text, max_cols,
                     cap_pt=11, cap_h=0.28):
    """Tile entries in a grid (max_cols=1 => vertical stack for wide images)."""
    slide = _new_slide(prs)
    add_textbox(slide, title, TITLE_LEFT, TITLE_TOP, TITLE_WIDTH, TITLE_HEIGHT,
                font_pt=title_font_for(title), color=BLACK, bold=True)
    n = len(entries)
    cols = min(max_cols, n)
    rows = math.ceil(n / cols)
    area_top = IMG_TOP
    area_h = FOOTER_TOP - IMG_TOP - 0.02
    cell_w = (SLIDE_W - 2 * MARGIN - (cols - 1) * GAP) / cols
    cell_h = (area_h - (rows - 1) * GAP) / rows
    img_h = cell_h - cap_h
    for i, (path, cap) in enumerate(entries):
        r, c = divmod(i, cols)
        left = MARGIN + c * (cell_w + GAP)
        top = area_top + r * (cell_h + GAP)
        add_image_in_box(slide, str(path), left, top, cell_w, img_h)
        _tile_caption(slide, cap, path, left, top + img_h, cell_w, cap_h, cap_pt)
    add_textbox(slide, footer_text, MARGIN, FOOTER_TOP, FOOTER_WIDTH,
                FOOTER_HEIGHT, font_pt=FOOTER_FONT_PT, color=GREY)
    return slide


def rel_footer(path):
    try:
        return Path(path).relative_to(ROOT).as_posix()
    except ValueError:
        return Path(path).name


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------
def _tiles(pairs):
    """(path, caption) pairs -> kept present ones; record missing."""
    kept, miss = [], []
    for path, cap in pairs:
        (kept if Path(path).exists() else miss).append((path, cap))
    return kept, [p for p, _ in miss]


def build_plan():
    """Hybrid layout across the 20 curated MT_loc_wrto_invag comparison figures:
    wide Mean±SEM profiles (02a-c) one-per-slide; enrichment / correlation /
    orientation / nuclear-shape violins grouped into per-section grids."""
    plan, missing = [], []

    # -----------------------------------------------------------------
    # 1. MT density vs NE geometry
    # -----------------------------------------------------------------
    plan.append(("divider", "MT density vs NE geometry",
                 "siCtrl vs DMSO · near the nuclear envelope"))

    # 02a-c -- Mean±SEM relative-density PROFILES (wide -> one per slide).
    prof_cap = "Mean±SEM relative density (÷ cell mean) · siCtrl vs DMSO"
    for fname, nice, extra in (
        ("02a_MT_avg_hulldist.png",
         "MT density vs invagination depth (hull-boundary distance)",
         " · x = distance below the convex hull (= invagination depth)"),
        ("02b_MT_avg_minCurvature.png", "MT density vs min curvature",
         " · concave (min curv < 0) toward the left"),
        ("02c_MT_avg_meanCurvature.png", "MT density vs mean curvature", ""),
    ):
        p = CURATED / fname
        if p.exists():
            plan.append(("single", nice, p, prof_cap + extra))
        else:
            missing.append(fname)

    # 03a-b -- enrichment in deep invaginations (2-up)
    kept, miss = _tiles([
        (CURATED / "03a_MT_enrichment_hulldist_gt0.5um.png", "hull dist > 0.5 μm"),
        (CURATED / "03b_MT_enrichment_hulldist_gt1.0um.png", "hull dist > 1.0 μm"),
    ])
    missing += miss
    if kept:
        plan.append(("grid", "MT enrichment in deep invaginations", kept,
                     PAIR_ENR, 2, 13, 0.30))

    # 03c-e -- enrichment on the concave surface (3-up)
    kept, miss = _tiles([
        (CURATED / "03c_MT_enrichment_minCurvature_lt0.png", "min curv < 0"),
        (CURATED / "03d_MT_enrichment_minCurvature_ltm0.25.png", "min curv < −0.25"),
        (CURATED / "03e_MT_enrichment_meanCurvature.png", "mean curv < 0"),
    ])
    missing += miss
    if kept:
        plan.append(("grid", "MT enrichment on the concave surface", kept,
                     PAIR_ENR, 3, 12, 0.30))

    # 04a-c -- per-cell correlation vs geometry (3-up)
    kept, miss = _tiles([
        (CURATED / "04a_MT_correlation_hulldist.png", "vs hull-boundary distance"),
        (CURATED / "04b_MT_correlation_minCurvature.png", "vs min curvature"),
        (CURATED / "04c_MT_correlation_meanCurvature.png", "vs mean curvature"),
    ])
    missing += miss
    if kept:
        plan.append(("grid", "MT per-cell correlation vs NE geometry", kept,
                     PAIR_CORR, 3, 12, 0.30))

    # -----------------------------------------------------------------
    # 2. MT orientation near invaginations
    # -----------------------------------------------------------------
    plan.append(("divider", "MT orientation near invaginations",
                 "siCtrl vs DMSO · orientation relative to the local NE normal"))

    p = CURATED / "06a_MT_avg_normalangle.png"
    if p.exists():
        plan.append(("single", "MT orientation vs the local NE normal", p,
                     "Mean±SEM angle between MT and the local NE normal · siCtrl vs DMSO"))
    else:
        missing.append(p.name)

    kept, miss = _tiles([
        (CURATED / "06b01_MT_orientation_ratio_above_below.png", "above ÷ below"),
        (CURATED / "06b02_MT_orientation_ratio_below_by_side.png", "below, by side"),
        (CURATED / "06b03_MT_orientation_ratio_above_by_side.png", "above, by side"),
    ])
    missing += miss
    if kept:
        plan.append(("grid", "MT orientation ratios", kept,
                     "siCtrl vs DMSO · one dot per cell", 3, 12, 0.30))

    # -----------------------------------------------------------------
    # 3. Nuclear shape
    # -----------------------------------------------------------------
    plan.append(("divider", "Nuclear shape",
                 "siCtrl vs DMSO · per-cell nuclear shape metrics"))
    kept, miss = _tiles([
        (CURATED / "07a_nuc_aspect_ratio.png", "aspect ratio"),
        (CURATED / "07b_nuc_height.png", "height (μm)"),
        (CURATED / "07c_nuc_solidity.png", "solidity"),
        (CURATED / "07d_chull_max_D.png", "convex-hull max diameter (μm)"),
    ])
    missing += miss
    if kept:
        plan.append(("grid", "Nuclear shape metrics", kept,
                     "siCtrl vs DMSO · one dot per cell", 2, 12, 0.30))

    # -----------------------------------------------------------------
    # 4. Cell counts
    # -----------------------------------------------------------------
    plan.append(("divider", "Cell counts",
                 "cells loaded → retained through each filter"))
    p = CURATED / "09a_cell_counts_filter_breakdown.png"
    if p.exists():
        plan.append(("single", "Cell-count filter breakdown", p,
                     "siCtrl (VimKD, n=107) vs DMSO (Noco, n=77) retained after filtering"))
    else:
        missing.append(p.name)

    return plan, missing


def main():
    list_only = "--list" in sys.argv
    plan, missing = build_plan()

    n_slides = 1 + len(plan)
    print("Compile: {}".format(ROOT))
    print("Output: {}".format(OUTPUT_PATH))
    print("{} content slides (+ title) = {} total\n".format(len(plan), n_slides))
    for it in plan:
        if it[0] == "divider":
            print("\n=== {} ===".format(it[1]))
        elif it[0] == "single":
            print("  [1] {}".format(it[1]))
        else:  # grid
            print("  [{}] {}".format(len(it[2]), it[1]))
    if missing:
        print("\nMISSING ({}): {}".format(len(missing), missing))
    if list_only:
        return

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    build_title_slide(prs, DECK_TITLE, DECK_SUBTITLE)
    for it in plan:
        if it[0] == "divider":
            build_divider_slide(prs, it[1], it[2])
        elif it[0] == "single":
            build_single_slide(prs, it[1], it[2], it[3])
        elif it[0] == "grid":
            _, title, entries, footer, max_cols, cap_pt, cap_h = it
            build_grid_slide(prs, title, entries, footer, max_cols, cap_pt, cap_h)

    if OUTPUT_PATH.exists():
        backup_dir = OUTPUT_PATH.parent / "backups"
        created = backup_presentation(str(OUTPUT_PATH), backup_base=str(backup_dir))
        if created:
            print("\nBacked up previous deck under: {}".format(backup_dir))

    prs.save(str(OUTPUT_PATH))
    total = len(prs.slides._sldIdLst)
    print("\nDone. {} slides written to:\n  {}".format(total, OUTPUT_PATH))
    if missing:
        print("Skipped {} missing panel(s).".format(len(missing)))


if __name__ == "__main__":
    main()
