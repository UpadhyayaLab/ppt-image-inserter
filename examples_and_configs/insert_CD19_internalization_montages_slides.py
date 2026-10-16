"""
insert_CD19_internalization_montages_slides.py

CAR-T CD19 internalization (Kiet, Mar 10, 2026 coculture), 10, 30 and 60 min.
CAT (left) vs FMC (right) side-by-side, ALL usable cells in numerical
order: every montage page pair gets two slides -
    1. CAR (cyan) / CD19 (magenta) merge
    2. CD19 grayscale + cell-body outline (cyan) + internalized puncta
       ringed yellow (the volscope MATLAB detector's kept objects)
Each tile label reads "cell<N>  z<plane>  n = <in plane> plane / <cell> total".

Montages are rendered by volscope dev/show_internalized_cells.m
(layout 'montage'), which writes per condition:
    <MONTAGE_ROOT>/<cond>/montages_<style>/montage_cells_<a>_<b>.png
CAT page k is paired with FMC page k (pages sorted by first cell number).
Layout, helpers and per-slide PPI (CAT and FMC at one physical scale)
are copied from insert_pMLC_puncta_montages_slides.py.

Usage:
    python examples_and_configs/insert_CD19_internalization_montages_slides.py
"""

import os
import re
import sys
from pathlib import Path
from typing import List, Optional, Tuple

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

EXP_LABEL = "Mar 10, 2026"
MONTAGE_ROOT = "K:/FF/volscope_show/kiet_FMC_internalized/png"
# (minutes, (left CAT folder, right FMC folder)) -- the configs/kiet_03102026.yaml folders
TIMEPOINTS = [
    (10, ("CAT_10min_001", "FMC_10min_")),
    (30, ("CAT_30min_",    "FMC_30min_")),
    (60, ("CAT_60min",     "FMC_60min001")),
]

# (montage style folder suffix, slide title with {tp}), in slide order per page pair.
KINDS = [
    ("merge",   "CAR and CD19, {tp} min, plane with most puncta"),
    ("cd19_hl", "Internalized CD19, {tp} min, plane with most puncta: outline, puncta"),
]

# Words of each title drawn in their channel color instead of naming the color
# (FF 2026-09-29); everything else in the title stays white.
TITLE_COLORS = {
    "merge":   {"CAR": "40E0FF", "CD19": "FF4FD8"},
    "cd19_hl": {"outline": "00FFFF", "puncta": "FFFF00"},
}


def color_title(slide, words):
    """Recolor the given whole words of the slide title (its first text box). With a
    colon in the title only the key after the LAST colon is colored, so 'puncta' in
    'plane with most puncta: outline, puncta' stays white."""
    tb = next(sh for sh in slide.shapes if sh.has_text_frame)
    para = tb.text_frame.paragraphs[0]
    text = para.text
    size, bold = para.runs[0].font.size, para.runs[0].font.bold
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    pat = re.compile(r"\b(" + "|".join(map(re.escape, words)) + r")\b")
    start = text.rfind(":") + 1                  # 0 when there is no colon
    segs, pos = [], 0
    for m in pat.finditer(text):
        if m.start() < start:
            continue
        segs += [(text[pos:m.start()], "FFFFFF"), (m.group(0), words[m.group(0)])]
        pos = m.end()
    segs.append((text[pos:], "FFFFFF"))
    for seg, col in segs:
        if not seg:
            continue
        run = para.add_run()
        run.text = seg
        run.font.size, run.font.bold = size, bold
        run.font.color.rgb = RGBColor.from_string(col)


# Channel key, upper left of every slide (FF 2026-09-29), in the image colors:
# CAR #40e0ff and CD19 #ff4fd8 as configs/kiet_03102026.yaml; outline/puncta as drawn.
LEGENDS = {
    "merge":   [("CAR", "40E0FF"), ("CD19", "FF4FD8")],
    "cd19_hl": [("CD19", "FFFFFF"), ("outline", "00FFFF"), ("puncta", "FFFF00")],
}


def add_legend(slide, entries, left=0.12, top=0.5, font_pt=14):
    """Channel names in their own colors, one line, upper left under the title."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(3.0), Inches(0.3))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.02)
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    for i, (name, hexcol) in enumerate(entries):
        run = para.add_run()
        run.text = name + ("   " if i < len(entries) - 1 else "")
        run.font.size = Pt(font_pt)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(hexcol)

OUTPUT_PATH = (
    "K:/FF/PPT/PPT_autogeneration/CART/CD19/"
    "CART_CD19_internalization_montages_10_30_60min_20260929.pptx"
)

CHUNK_GLOB = "montage_cells_*.png"

# Colors
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)

SLIDE_W = 13.333
SLIDE_H = 7.5

TITLE_LEFT = 0.05
TITLE_TOP = 0.05
TITLE_WIDTH = SLIDE_W - 2 * 0.05
TITLE_HEIGHT = 0.40
TITLE_FONT_PT = 24

GRID_LEFT = 0.05
GRID_TOP = 0.50
CELL_W = 6.60
CELL_H = SLIDE_H - GRID_TOP - 0.05
LABEL_H = 0.22
IMG_H = CELL_H - LABEL_H
LABEL_FONT_PT = 14
COL_GAP = SLIDE_W - 2 * GRID_LEFT - 2 * CELL_W

CELL_POSITIONS = [
    (GRID_LEFT,                    GRID_TOP),
    (GRID_LEFT + CELL_W + COL_GAP, GRID_TOP),
]


def _long_path_str(p: Path) -> str:
    s = str(p)
    if os.name == "nt" and len(s) >= 240 and not s.startswith("\\\\?\\"):
        s = s.replace("/", "\\")
        if not s.startswith("\\\\"):
            return "\\\\?\\" + s
    return s


def _parse_chunk_range(p: Path) -> Tuple[int, int]:
    m4 = re.match(r"montage_cells_(\d+)_(\d+)_(\d+)_(\d+)\.png$", p.name)
    if m4:
        f_a, c_a, f_b, c_b = (int(x) for x in m4.groups())
        return (f_a * 1000 + c_a, f_b * 1000 + c_b)
    m2 = re.match(r"montage_cells_(\d+)_(\d+)\.png$", p.name)
    if m2:
        return (int(m2.group(1)), int(m2.group(2)))
    return (0, 0)


def _list_chunk_files(montages_dir):
    """Long-path-safe list of montage chunk PNGs (unsorted); [] if dir absent.
    pathlib .is_dir()/.glob() silently return False/empty past Windows MAX_PATH
    (260) even when the dir exists, so enumerate via os.listdir over the
    \\?\-prefixed path."""
    import fnmatch
    long_dir = _long_path_str(montages_dir)
    if not os.path.isdir(long_dir):
        return []
    return [montages_dir / n for n in os.listdir(long_dir)
            if fnmatch.fnmatch(n, CHUNK_GLOB)]


def _png_dims(path: Path) -> Tuple[int, int]:
    with Image.open(_long_path_str(path)) as im:
        return im.size


def _path_exists(p: Path) -> bool:
    return os.path.exists(_long_path_str(p))


def compute_deck_ppi(image_paths: List[Path], max_w_in: float, max_h_in: float) -> float:
    ppi = 0.0
    for p in image_paths:
        w_px, h_px = _png_dims(p)
        ppi = max(ppi, w_px / max_w_in, h_px / max_h_in)
    return ppi


def add_textbox(slide, text, left, top, width, height, font_pt, color, bold=False):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.text = text
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.runs[0]
    run.font.size = Pt(font_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def set_slide_background(slide, rgb: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_image_in_cell_at_ppi(slide, image_path: Path, ppi: float,
                             cell_left: float, cell_top: float):
    w_px, h_px = _png_dims(image_path)
    w_in = w_px / ppi
    h_in = h_px / ppi
    img_area_top = cell_top + LABEL_H
    left_in = cell_left + (CELL_W - w_in) / 2
    top_in  = img_area_top + (IMG_H - h_in) / 2
    return slide.shapes.add_picture(
        _long_path_str(image_path),
        Inches(left_in),
        Inches(top_in),
        width=Inches(w_in),
    )


def build_compare_slide(prs, title_text: str,
                        cat_img: Optional[Path], fmc_img: Optional[Path],
                        deck_ppi: float, labels: Tuple[str, str] = ("CAT", "FMC"),
                        empty_text: str = "(missing)"):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_slide_background(slide, BLACK)

    add_textbox(
        slide, title_text,
        TITLE_LEFT, TITLE_TOP, TITLE_WIDTH, TITLE_HEIGHT,
        font_pt=TITLE_FONT_PT, color=WHITE, bold=True,
    )

    cells = [
        (labels[0], cat_img, CELL_POSITIONS[0]),
        (labels[1], fmc_img, CELL_POSITIONS[1]),
    ]
    missing = []
    for label, img_path, (cell_left, cell_top) in cells:
        add_textbox(
            slide, label,
            cell_left, cell_top, CELL_W, LABEL_H,
            font_pt=LABEL_FONT_PT, color=WHITE, bold=True,
        )
        if img_path is not None and _path_exists(img_path):
            add_image_in_cell_at_ppi(slide, img_path, deck_ppi, cell_left, cell_top)
        else:
            add_textbox(
                slide, empty_text,
                cell_left, cell_top + LABEL_H + IMG_H / 2 - 0.15, CELL_W, 0.3,
                font_pt=14, color=WHITE,
            )
            missing.append(label)
    return slide, missing



def list_pages(cond: str, style: str) -> List[Path]:
    """This condition's numbered montage pages, sorted by first cell number."""
    pages = _list_chunk_files(Path(MONTAGE_ROOT) / cond / f"montages_{style}")
    return sorted(pages, key=lambda p: _parse_chunk_range(p)[0])


def page_label(cond: str, page: Optional[Path]) -> str:
    cell = cond.split("_")[0]
    if page is None:
        return cell
    a, b = _parse_chunk_range(page)
    return f"{cell} (cells {a}–{b})"


def main() -> None:
    out_path = Path(OUTPUT_PATH)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    missing_total: List[str] = []
    for tp, conds in TIMEPOINTS:
        pages = {(c, s): list_pages(c, s) for c in conds for s, _ in KINDS}
        for (c, s), p in pages.items():
            print(f"{c} {s}: {len(p)} pages  {[q.name for q in p]}")
        n_pairs = max(len(p) for p in pages.values())
        for k in range(n_pairs):
            for style, subtitle in KINDS:
                imgs = [pages[(c, style)][k] if k < len(pages[(c, style)]) else None
                        for c in conds]
                present = [p for p in imgs if p is not None and _path_exists(p)]
                slide_ppi = compute_deck_ppi(present, CELL_W, IMG_H) if present else 100.0
                # FF 2026-09-29: compact, no dashes, say it is one plane
                title = subtitle.format(tp=tp)
                labels = tuple(page_label(c, p) for c, p in zip(conds, imgs))
                slide, missing = build_compare_slide(prs, title, imgs[0], imgs[1],
                                                 slide_ppi, labels,
                                                 "(all cells shown on earlier slides)")
                add_legend(slide, LEGENDS[style])
                color_title(slide, TITLE_COLORS[style])
                print(f"[{tp} min page {k + 1} {style}]  PPI={slide_ppi:7.2f}  "
                      f"{[p.name if p else None for p in imgs]}")
                missing_total += [f"{tp} min page {k + 1} {style} {m}" for m in missing]

    prs.save(str(out_path))
    print(f"\nDone. {len(prs.slides)} slides written to:\n  {out_path}")
    if missing_total:
        print(f"\nMissing ({len(missing_total)}):")
        for m in missing_total:
            print(f"  - {m}")
    else:
        print("\nAll images found - no missing items.")


if __name__ == "__main__":
    main()
